"""대용량 엑셀 작성기: 파일 형식, 시트 분할, 값/서식, 특수문자."""
import zipfile

from openpyxl import load_workbook

from app.xlsx_stream import XlsxStreamWriter, _col

HEAD = ["판매년월", "매장코드", "수량", "금액"]
KINDS = ["text", "text", "int", "int"]
WIDTHS = [10, 10, 8, 12]


def _read(path):
    wb = load_workbook(path, read_only=True)
    try:
        return wb.sheetnames, [list(ws.iter_rows(values_only=True)) for ws in wb.worksheets]
    finally:
        wb.close()


def test_column_letters():
    assert [_col(i) for i in (0, 25, 26, 27, 51, 52, 701, 702)] == ["A", "Z", "AA", "AB", "AZ", "BA", "ZZ", "AAA"]


def test_rows_split_into_sheets(tmp_path):
    path = tmp_path / "t.xlsx"
    rows = [(f"2026{m:02d}", f"S{i:05d}", i, i * 1000) for m in (1, 2) for i in range(4)]  # 8행
    with XlsxStreamWriter(str(path), HEAD, KINDS, WIDTHS, sheet_name="판매집계", sheet_rows=3) as w:
        w.write_rows(rows[:5])
        w.write_rows(rows[5:])  # 여러 번 나눠 써도 이어진다
    assert w.total == 8
    names, sheets = _read(path)
    assert names == ["판매집계", "판매집계_2", "판매집계_3"]
    assert [len(s) - 1 for s in sheets] == [3, 3, 2]
    assert all(tuple(s[0]) == tuple(HEAD) for s in sheets)  # 시트마다 헤더
    data = [r for s in sheets for r in s[1:]]
    assert data == [tuple(r) for r in rows]  # 순서·값 보존


def test_values_none_and_special_chars(tmp_path):
    path = tmp_path / "t.xlsx"
    rows = [
        ("202608", None, 1, None),                     # 빈 칸이 있어도 열 위치가 밀리지 않음
        ("<a&b>\"'", "제어\x01문자\x0b제거", -4, 12.5),  # XML 특수문자 · 엑셀 불가 제어문자
        ("", "매장", 0, 10**12),
    ]
    with XlsxStreamWriter(str(path), HEAD, KINDS, WIDTHS) as w:
        w.write_rows(rows)
    _, sheets = _read(path)
    body = sheets[0][1:]
    assert body[0] == ("202608", None, 1, None)
    assert body[1] == ("<a&b>\"'", "제어문자제거", -4, 12.5)
    assert body[2] == (None, "매장", 0, 10**12)


def test_numbers_have_thousands_format_and_frozen_header(tmp_path):
    path = tmp_path / "t.xlsx"
    with XlsxStreamWriter(str(path), HEAD, KINDS, WIDTHS) as w:
        w.write_rows([("202608", "S1", 1234, 5678)])
    with zipfile.ZipFile(path) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        styles = z.read("xl/styles.xml").decode()
    assert 'state="frozen"' in sheet
    assert '<c s="2"><v>1234</v></c>' in sheet
    assert 'numFmtId="3"' in styles  # #,##0


def test_empty_result_still_valid(tmp_path):
    path = tmp_path / "t.xlsx"
    with XlsxStreamWriter(str(path), HEAD, KINDS, WIDTHS) as w:
        w.write_rows([])
    names, sheets = _read(path)
    assert names == ["Sheet"] and [tuple(r) for r in sheets[0]] == [tuple(HEAD)]


def test_exception_does_not_leave_partial_close(tmp_path):
    path = tmp_path / "t.xlsx"
    try:
        with XlsxStreamWriter(str(path), HEAD, KINDS, WIDTHS) as w:
            w.write_rows([("202608", "S1", 1, 1)])
            raise RuntimeError("중단")
    except RuntimeError:
        pass
    assert w._f is None or w._f.closed
