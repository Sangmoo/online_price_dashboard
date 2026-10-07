"""엑셀 업로드(실사계획 · 판매처 매장 연결) 미리보기 · 저장, 쿼리 성능 통계 · 로그 분석 · 실행 계획 검증. DB 는 흉내 낸다."""
import base64
import io
from datetime import datetime

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app import invt_plan, logs, mall_shop, sql_perf, sql_trace, uploads


def _file(template: bytes, rows: list[list]) -> dict:
    wb = load_workbook(io.BytesIO(template))
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return {"file": "data:application/octet-stream;base64," + base64.b64encode(buf.getvalue()).decode()}


@pytest.fixture
def invt_db(monkeypatch):
    shops = {"A11001": "현대서울", "T22002": "롯데본점"}

    def chunks(sql, ids):
        if "COUNT(*)" in sql:
            return [("A11001", 2)] if "A11001" in ids else []
        return [(i,) for i in ids if i in shops]

    def detail(sid):
        return {"values": {"shopId": sid, "shopNm": shops[sid], "brdNm": "시스티나", "shopRankNm": "A", "stockQty": 100},
                "missing": ["prevSaleAmt"], "notes": []}

    monkeypatch.setattr(uploads, "_in_chunks", chunks)
    monkeypatch.setattr(invt_plan, "shop_detail", detail)
    return shops


def test_invt_preview_autofills_and_flags(invt_db):
    body = _file(uploads.invt_template(), [
        ["A11001", datetime(2026, 11, 3), "정기", "300,000", None, "1", "Y", None, "경기도"],
        ["t22002", "미정", None, None, None, None, None, None, None, None, None, None, "비고"],
        ["ZZ0001", "2026-11-04"],
        ["A11001", "2026-13-40"],
        [None, None, None, None],          # 빈 행은 건너뜀
    ])
    res = uploads.invt_preview(body)
    assert res["summary"] == {"total": 4, "ok": 1, "warn": 1, "error": 2}
    a, t, z, bad = res["rows"]
    assert a["status"] == "warn" and any("이미 등록된 실사계획 2건" in m for m in a["messages"])
    assert a["values"]["shopNm"] == "현대서울" and a["values"]["shopRankNm"] == "A"              # 매장코드로 자동 입력
    assert a["values"]["regionNm"] == "수도권" and a["values"]["stlmTeam"] == "1팀" and a["values"]["invtPlanDt"] == "20261103"
    assert t["status"] == "ok" and t["values"]["shopId"] == "T22002" and "invtPlanDt" not in t["values"]
    assert any("전년 매출" in m for m in t["messages"])                                          # 자동으로 못 채운 항목 안내
    assert z["status"] == "error" and "ZZ0001" in z["messages"][0]
    assert bad["status"] == "error" and "4행" not in bad["messages"][0]


def test_invt_apply_inserts_valid_rows_in_one_transaction(invt_db, monkeypatch):
    calls = {}

    class Cur:
        def executemany(self, sql, params):
            calls["sql"], calls["params"] = sql, params

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    class Conn(Cur):
        def cursor(self):
            return Cur()

        def commit(self):
            calls["commit"] = True

    class Pool:
        def acquire(self):
            return Conn()

    monkeypatch.setattr(uploads.db, "get_pool", lambda: Pool())
    res = uploads.invt_apply({"id": "900001"}, {"rows": [
        {"row": 2, "values": {"shopId": "A11001", "invtPlanDt": "20261103", "shopNm": "엑셀에 적은 이름"}},
        {"row": 3, "values": {"shopId": "ZZ0001"}},
    ]})
    assert res == {"saved": 1, "skipped": 1, "errors": [{"row": 3, "messages": ["매장(T_SHOP)에 없는 매장코드: ZZ0001"]}]}
    assert "SQ_SHOP_INVT_PLAN.NEXTVAL" in calls["sql"] and calls["commit"]
    p = calls["params"][0]
    assert p["SHOP_NM"] == "엑셀에 적은 이름" and p["SHOP_RANK_NM"] == "A" and p["INS_USERID"] == "900001"   # 직접 적은 값 우선


def test_upload_rejects_bad_files():
    with pytest.raises(HTTPException):
        uploads.invt_preview({"file": base64.b64encode(b"not excel").decode()})
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["아무", "머리글"])
    wb.active.append(["x", "y"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(HTTPException) as ex:
        uploads.invt_preview({"file": base64.b64encode(buf.getvalue()).decode()})
    assert "매장코드" in ex.value.detail["message"]


@pytest.fixture
def mall_db(monkeypatch):
    monkeypatch.setattr(mall_shop, "table_ready", lambda: True)
    monkeypatch.setattr(mall_shop, "_shop_names", lambda ids: {i: f"매장{i}" for i in ids if i.startswith("S")})
    monkeypatch.setattr(mall_shop, "_maps", lambda: {
        ("하프클럽", "-", "S"): {"SHOP_ID": "S51005", "USE_YN": "Y", "RMK": "기존 비고"},
        ("롯데온", "123", "T"): {"SHOP_ID": "S00001", "USE_YN": "Y", "RMK": None},
    })


def test_mall_preview_changes_and_keeps_blank_fields(mall_db):
    body = _file(uploads.mall_template(), [
        ["하프클럽", None, "쉬즈미스", "S51005"],           # 같음 (비고 비움 → 기존 유지)
        ["하프클럽", None, "S", "S51006", "N"],            # 같은 키 중복 → 오류
        ["롯데온", "123", "T", "S00002"],                  # 변경
        ["새몰", None, None, "S1"],                        # 신규 (브랜드 비움 = *)
        ["롯데온", "123", "리스트", None],                  # 같은 키 중복
        ["없는몰", None, "A", None],                       # 해제할 것 없음 → 주의
        ["몰", None, "Q", "S1"],                           # 브랜드 오류
        ["몰2", None, "S", "X9"],                          # 없는 매장
    ])
    res = uploads.mall_preview(body)
    rows = {r["row"]: r for r in res["rows"]}
    assert rows[2]["change"] == "같음" and rows[2]["values"]["rmk"] == "기존 비고"
    assert rows[3]["status"] == "error" and rows[6]["status"] == "error"
    assert rows[4]["change"] == "변경" and "S00001 → S00002" in rows[4]["messages"][0]
    assert rows[5]["change"] == "신규" and rows[5]["values"]["brdCd"] == "*"
    assert rows[7]["status"] == "warn"
    assert rows[8]["status"] == "error" and "브랜드" in rows[8]["messages"][0]
    assert rows[9]["status"] == "error" and "X9" in rows[9]["messages"][0]
    assert res["summary"]["changes"] == {"신규": 1, "변경": 1, "해제": 0, "같음": 1}


def test_mall_apply_saves_only_changes(mall_db, monkeypatch):
    saved = {}
    monkeypatch.setattr(mall_shop, "save", lambda admin, items: saved.setdefault("items", items) and {"saved": len(items), "deleted": 0, "changed": len(items)})
    res = uploads.mall_apply({"id": "900001"}, {"rows": [
        {"row": 2, "values": {"mallNm": "하프클럽", "sellNo": "-", "brdCd": "S", "shopId": "S51005", "useYn": "Y", "rmk": "기존 비고"}},
        {"row": 3, "values": {"mallNm": "새몰", "sellNo": "-", "brdCd": "*", "shopId": "S1", "useYn": "Y", "rmk": None}},
    ]})
    assert [i["mallNm"] for i in saved["items"]] == ["새몰"] and res["skipped"] == 1


# ---------------------------------------------------------------------------- 쿼리 성능
def test_perf_stats_and_logs(fake_oracle, tmp_path, monkeypatch):
    from app import roles

    sql_trace.clear()
    roles._rows()
    roles._rows()
    log = tmp_path / "app.log"
    today = datetime.now().strftime("%Y-%m-%d")
    log.write_text("\n".join([
        f"{today} 09:00:00 INFO  erp.request  user=1 GET /api/sale-monthly 200 1200ms",
        f"{today} 09:00:01 WARNING erp.request  user=1 GET /api/sale-monthly 200 9000ms",
        f"{today} 09:00:02 INFO  erp.request  user=1 GET /api/invt-plans/shops/A11001 200 100ms",
        f"{today} 09:00:03 WARNING erp.sql      느린 query 4.50s | binds=['m0'] | SELECT * FROM T WHERE A = :m0",
        "2000-01-01 09:00:00 INFO  erp.request  user=1 GET /api/old 200 1ms",
    ]), encoding="utf-8")
    monkeypatch.setattr(logs, "LOG_DIR", tmp_path)
    sql_perf._log_cache.clear()
    r = sql_perf.report(7)
    fn = next(f for f in r["funcs"] if f["fn"] == "roles._rows")
    assert fn["count"] == 2 and fn["where"] == [{"menu": "관리자", "feature": "권한 묶음"}]
    assert r["slowSql"][0]["raw"]
    req = {x["path"]: x for x in r["requests"]}
    assert req["/api/sale-monthly"]["count"] == 2 and req["/api/sale-monthly"]["slow"] == 1 and req["/api/sale-monthly"]["maxMs"] == 9000
    assert "/api/invt-plans/shops/{id}" in req and "/api/old" not in req
    assert req["/api/sale-monthly"]["menu"] == "월별 매장별 판매 집계"
    assert r["logSlowSql"][0]["maxMs"] == 4500 and r["daily"][0]["slow"] == 1
    sql_perf.clear()
    assert sql_perf.report(7)["funcs"] == []


def test_explain_rejects_unsafe_sql():
    for bad in ("DROP TABLE X", "SELECT 1 FROM DUAL; DELETE FROM X", "SELECT {cols} FROM T", "SELECT A FROM T WHERE …"):
        with pytest.raises(HTTPException):
            sql_perf._check(bad)
    assert sql_perf._check("SELECT ';' FROM DUAL;") == "SELECT ';' FROM DUAL"      # 문자열 안 ; 는 괜찮음, 끝 ; 는 뗌
    assert sql_perf.sql_id("SELECT COUNT(*) FROM V$SQL WHERE ROWNUM = 1") == "bxg5rxdw1gwfs"
