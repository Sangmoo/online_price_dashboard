"""판매 현황 계산: 기간·비교 기준 해석, 목표 배분(매장 단위 · 판매 브랜드 기준), 브랜드 조건, 매장 순위, 보고용 엑셀. DB 는 가짜로 대체."""
import io

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app import sale_dashboard as sd
from app import sale_dashboard_report as rpt

M = 1_000_000
MONTHS = [f"2026{m:02d}" for m in range(8, 0, -1)] + [f"2025{m:02d}" for m in range(12, 0, -1)]
TEAMS = ["쉬즈1팀", "쉬즈2팀", "쉬즈4팀", "리스트1팀", "리스트2팀"]
# (월, 매장, 팀, 실판금액) — 수량·할인·원가는 금액에서 파생
SALES = [
    ("202608", "S1", "쉬즈1팀", 100 * M), ("202508", "S1", "쉬즈1팀", 80 * M), ("202607", "S1", "쉬즈1팀", 90 * M),
    ("202608", "S2", "쉬즈4팀", 50 * M), ("202508", "S2", "쉬즈4팀", 60 * M),   # 여러 브랜드 목표가 있는 매장
    ("202608", "S3", "리스트1팀", 40 * M),                                        # 신규 · 목표 없음
    ("202608", "S5", "리스트2팀", 70 * M), ("202508", "S5", "리스트2팀", 65 * M),
    ("202608", "S6", "쉬즈2팀", 5 * M), ("202508", "S6", "쉬즈2팀", 30 * M),     # 폐점 표시 매장
]
GOALS = [("S1", "S", 120 * M), ("S2", "S", 30 * M), ("S2", "T", 20 * M), ("S2", "A", 10 * M), ("S4", "T", 25 * M), ("S5", "T", 100 * M)]
NAMES = {"S1": "매장1", "S2": "복합매장", "S3": "신규매장", "S4": "미오픈매장", "S5": "매장5", "S6": "매장6(폐)"}


@pytest.fixture
def fake_db(monkeypatch):
    seen = []

    def query(sql, params=None):
        p = params or {}
        seen.append(sql)
        months = {v for k, v in p.items() if k.startswith("m")}
        teams = {v for k, v in p.items() if k.startswith("t")}
        if "SELECT DISTINCT MAKE_YYMM" in sql:
            return ["M"], [(m,) for m in MONTHS]
        if "SELECT DISTINCT TEAM_CD" in sql:
            return ["T"], [(t,) for t in TEAMS]
        if "T_SHOP_SELL_MGOAL" in sql:
            return ["S", "B", "G"], [g for g in GOALS] if "202608" in months else []
        if "GROUP BY MAKE_YYMM, SHOP_ID, TEAM_CD" in sql:
            return [], [(m, s, t, a, a // M, a // 10, a * 3 // 10) for m, s, t, a in SALES if m in months]
        if "GROUP BY MAKE_YYMM" in sql:
            out = {}
            for m, s, t, a in SALES:
                if m in months and (not teams or t in teams):
                    o = out.setdefault(m, [m, 0, 0, 0, 0, set()])
                    o[1] += a
                    o[2] += a // M
                    o[3] += a // 10
                    o[4] += a * 3 // 10
                    o[5].add(s)
            return [], [(o[0], o[1], o[2], o[3], o[4], len(o[5])) for o in out.values()]
        raise AssertionError(sql)

    monkeypatch.setattr(sd.db, "query", query)
    monkeypatch.setattr(sd, "shop_names", lambda ids: {i: NAMES[i] for i in ids if i in NAMES})
    sd.clear_cache()
    yield seen
    sd.clear_cache()


def test_resolve_periods(fake_db):
    r = sd.resolve()
    assert (r["from"], r["to"], r["base"], r["extra"], r["extraLabel"]) == ("202608", "202608", ["202508"], ["202607"], "전월")
    r = sd.resolve("202608", "202601")
    assert r["base"] == [f"2025{m:02d}" for m in range(1, 9)] and r["extra"][0] == "202505" and r["extraLabel"] == "직전 기간"
    r = sd.resolve("202608", "202601", "prev")
    assert r["base"] == ["202505", "202506", "202507", "202508", "202509", "202510", "202511", "202512"] and r["extraLabel"] == "전년 동기"
    r = sd.resolve("2026-08", None, "custom", "202501", "202503", "쉬즈미스")
    assert r["base"] == ["202501", "202502", "202503"] and r["teams"] == ["쉬즈1팀", "쉬즈2팀", "쉬즈4팀"]
    assert sd.resolve(brand="쉬즈")["brand"] == "쉬즈미스"  # 예전 이름으로 찾아도 같은 브랜드
    assert r["ytd"][0] == "202601" and r["ytdPrev"][-1] == "202508"


@pytest.mark.parametrize("kw,msg", [
    ({"ym": "202608", "frm": "202501"}, "최대 12개월"),
    ({"ym": "202601", "frm": "202608"}, "늦습니다"),
    ({"ym": "209901"}, "판매 데이터가 없습니다"),
    ({"ym": "2026-13"}, "YYYYMM"),
    ({"cmp": "week"}, "비교 기준"),
    ({"cmp": "custom"}, "비교 기간"),
    ({"brand": "없는브랜드"}, "브랜드"),
])
def test_resolve_errors(fake_db, kw, msg):
    with pytest.raises(HTTPException) as ex:
        sd.resolve(**kw)
    assert msg in ex.value.detail["message"]


def test_goals_are_attributed_to_the_shops_sales_brand(fake_db):
    d = sd.dashboard(full=True)
    k = d["kpi"]
    assert (k["amt"], k["baseAmt"], k["extraAmt"]) == (265 * M, 235 * M, 90 * M)
    # 목표: S1 120 + S2(S·T·A 합) 60 + S4(판매 없음) 25 + S5 100 = 305 / 목표 매장 매출 = 100 + 50 + 70
    assert (k["goalAmt"], k["goalSalesAmt"], k["achieve"], k["goalShops"]) == (305 * M, 220 * M, 72.1, 4)
    assert (k["noGoalShops"], k["noGoalAmt"]) == (2, 45 * M)
    brands = {b["brand"]: b for b in d["brands"]}
    assert (brands["쉬즈미스"]["amt"], brands["쉬즈미스"]["goalAmt"], brands["쉬즈미스"]["achieve"]) == (155 * M, 180 * M, 83.3)  # 복합매장 목표 전부 쉬즈미스로
    assert (brands["리스트"]["goalAmt"], brands["리스트"]["achieve"]) == (125 * M, 56.0)  # 판매 없는 S4 는 목표 브랜드코드 T → 리스트
    assert brands["쉬즈미스"]["share"] == 58.5 and set(brands) == {"쉬즈미스", "리스트"}
    assert d["brandOptions"] == ["리스트", "쉬즈미스"] and {s["brand"] for s in d["allShops"]} == {"쉬즈미스", "리스트"}
    assert d["hasGoals"] and len(d["allShops"]) == 6


def test_rankings_exclude_closed_and_small_shops(fake_db):
    d = sd.dashboard()
    assert "allShops" not in d  # 화면 응답에는 전체 매장 목록을 싣지 않는다
    assert [s["shopId"] for s in d["topShops"]] == ["S1", "S5", "S2", "S3", "S6"]
    assert [s["shopId"] for s in d["risers"]] == ["S1", "S5", "S2"]
    assert [s["shopId"] for s in d["fallers"]] == ["S2", "S5", "S1"]
    assert d["shopCounts"] == {"selling": 5, "new": 1, "comparable": 3, "closedExcluded": 1}
    assert [s["shopId"] for s in d["laggards"]] == ["S5", "S1", "S2"]  # 70%, 83.3%(목표 120), 83.3%(목표 60): 같으면 목표 큰 순
    assert d["risers"][0]["change"] == 25.0 and d["topShops"][0]["achieve"] == 83.3


def test_brand_filter(fake_db):
    d = sd.dashboard(brand="리스트")
    k = d["kpi"]
    assert (k["amt"], k["goalAmt"], k["goalSalesAmt"]) == (110 * M, 125 * M, 70 * M)
    assert [b["brand"] for b in d["brands"]] == ["리스트"] and d["brand"] == "리스트"
    assert d["trend"][-1]["amt"] == 110 * M  # 추이도 브랜드 팀으로 거른 값
    assert all(t["brand"] == "리스트" for t in d["teams"])


def test_cache_and_clear(fake_db):
    sd.dashboard()
    n = len(fake_db)
    sd.dashboard()
    assert len(fake_db) == n  # 같은 조건은 캐시
    sd.clear_cache()
    sd.dashboard()
    assert len(fake_db) > n


def test_report_workbook(fake_db):
    d = sd.dashboard(ym="202608", cmp="prev", full=True)
    wb = load_workbook(io.BytesIO(rpt.build(d)))
    assert wb.sheetnames == ["요약", "월별 추이", "브랜드별", "팀별", "매장 순위", "전체 매장"]
    assert "비교 직전 기간 2026-07" in wb["요약"]["A2"].value
    shops = wb["전체 매장"]
    assert [c.value for c in shops[4]][:3] == ["순위", "매장명", "매장코드"] and shops.max_row == 4 + 6
    assert rpt.filename(d) == "판매현황_202608_prev.xlsx"


def test_brand_of():
    assert [sd.brand_of(t) for t in ("쉬즈3팀", "리스트1팀", "시스티나5팀", None)] == ["쉬즈미스", "리스트", "시스티나", "(미지정)"]
