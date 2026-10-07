"""엑셀 업로드로 한 번에 등록: 매장 재고 실사계획 · 판매처 매장 연결.

흐름: [양식 내려받기] → 작성 → 업로드(미리보기: 행마다 정상 · 주의 · 오류) → [정상 n건 저장].
- 미리보기는 저장하지 않는다. 저장할 때 서버가 다시 검증하고 오류 행은 빼고 저장한다.
- 실사계획: 매장코드만 있으면 화면에서 매장을 고를 때처럼 매장 정보(브랜드 · 유통 · 매장명 · 매출 · 최종실사 · 재고 · 관리등급 ·
  매장번호)를 자동으로 채운다. 엑셀에 같은 항목을 직접 적으면 그 값을 쓴다. 한 번에 한 트랜잭션으로 등록.
- 판매처 매장 연결: 기존 저장(mall_shop.save)을 그대로 써서 MERGE · 변경 이력.
"""
from __future__ import annotations

import base64
import io
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import db, invt_plan, mall_shop
from .shop_info import BRAND_CODES

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = {"invt": 300, "mall": 500}
HEAD_FILL = PatternFill("solid", fgColor="1F1F7A")
REQ_FILL = PatternFill("solid", fgColor="B91C1C")


def _bad(msg: str):
    raise HTTPException(400, {"message": msg, "code": "BAD_REQUEST"})


def _key(label: str) -> str:
    """머리글 비교용: 공백 · * · 괄호 설명 제거"""
    return re.sub(r"\(.*?\)|[\s*]", "", str(label or ""))


# ----------------------------------------------------------------------------
# 공통: 양식 만들기 · 읽기
# ----------------------------------------------------------------------------
def _template(title: str, cols: list[tuple[str, str, bool, int]], guide: list[tuple[str, str]],
              lists: dict[str, list[str]] | None = None) -> bytes:
    """cols: (키, 머리글, 필수, 너비). lists: 키 → 선택 목록 (셀 드롭다운)"""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for ci, (key, label, req, width) in enumerate(cols, start=1):
        c = ws.cell(1, ci, f"{label}*" if req else label)
        c.fill = REQ_FILL if req else HEAD_FILL
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width = width
        if lists and key in lists:
            dv = DataValidation(type="list", formula1='"' + ",".join(lists[key]) + '"', allow_blank=True)
            dv.error, dv.errorTitle = "목록에서 고르세요.", "입력 값 확인"
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(ci)}2:{get_column_letter(ci)}1000")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    g = wb.create_sheet("작성 안내")
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 100
    for ri, (a, b) in enumerate(guide, start=1):
        g.cell(ri, 1, a).font = Font(bold=True)
        g.cell(ri, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _decode(body: dict) -> bytes:
    raw = str((body or {}).get("file") or "")
    if "," in raw[:100] and raw.startswith("data:"):
        raw = raw.split(",", 1)[1]
    try:
        data = base64.b64decode(raw, validate=False)
    except ValueError:
        _bad("파일을 읽을 수 없습니다.")
    if not data:
        _bad("엑셀 파일을 선택하세요.")
    if len(data) > MAX_BYTES:
        _bad("파일이 너무 큽니다 (최대 5MB).")
    if not data.startswith(b"PK"):
        _bad("엑셀(.xlsx) 파일만 올릴 수 있습니다. (.xls 는 엑셀에서 .xlsx 로 다시 저장하세요)")
    return data


def _cell(v):
    if isinstance(v, datetime):
        return v.strftime("%Y%m%d")
    if isinstance(v, date):
        return v.strftime("%Y%m%d")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _read(data: bytes, headers: dict[str, str], required: list[str], limit: int) -> list[tuple[int, dict]]:
    """첫 시트에서 머리글 행(처음 10행 안)을 찾아 [(엑셀 행 번호, {키: 값})]. 빈 행은 건너뛴다."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        _bad("엑셀 파일을 열 수 없습니다. 양식 파일에 작성했는지 확인하세요.")
    ws = wb.worksheets[0]
    lookup = {_key(label): key for label, key in headers.items()}
    rows = ws.iter_rows(values_only=True)
    colmap: dict[int, str] = {}
    head_no = 0
    for no, row in enumerate(rows, start=1):
        found = {i: lookup[_key(v)] for i, v in enumerate(row) if v is not None and _key(v) in lookup}
        if found:
            colmap, head_no = found, no
            break
        if no >= 10:
            break
    inv = {v: k for k, v in headers.items()}
    missing = [inv[k] for k in required if k not in colmap.values()]
    if not colmap or missing:
        _bad(f"양식의 머리글을 찾을 수 없습니다. 필수 열: {', '.join(missing or [inv[k] for k in required])} — [양식 내려받기] 파일을 쓰세요.")
    out = []
    for no, row in enumerate(rows, start=head_no + 1):
        vals = {key: _cell(row[i]) for i, key in colmap.items() if i < len(row)}
        if all(v is None for v in vals.values()):
            continue
        out.append((no, vals))
        if len(out) > limit:
            _bad(f"한 번에 {limit}행까지 올릴 수 있습니다. 나눠서 올리세요.")
    wb.close()
    if not out:
        _bad("입력된 행이 없습니다.")
    return out


def _msg(ex: Exception) -> str:
    if isinstance(ex, HTTPException) and isinstance(ex.detail, dict):
        return str(ex.detail.get("message"))
    return str(ex).splitlines()[0][:200]


def _summary(rows: list[dict]) -> dict:
    return {"total": len(rows), "ok": sum(r["status"] == "ok" for r in rows), "warn": sum(r["status"] == "warn" for r in rows),
            "error": sum(r["status"] == "error" for r in rows)}


def _in_chunks(sql: str, ids: list[str]) -> list[tuple]:
    out: list[tuple] = []
    for i in range(0, len(ids), 900):
        part = ids[i:i + 900]
        binds = {f"i{n}": v for n, v in enumerate(part)}
        out += db.query(sql.format(ids=", ".join(":" + k for k in binds)), binds)[1]
    return out


# ----------------------------------------------------------------------------
# 매장 재고 실사계획
# ----------------------------------------------------------------------------
INVT_COLS = [  # (키, 머리글, 필수, 너비) — 매장 정보는 매장코드로 자동 입력
    ("shopId", "매장코드", True, 12), ("invtPlanDt", "실사예정일", False, 13), ("invtPlanNote", "실사예정", False, 24),
    ("baseFee", "기본료", False, 12), ("expectAmt", "실사예상액", False, 13), ("stlmTeam", "정산 팀구분", False, 12),
    ("twiceYearYn", "연2회 실사 매장", False, 14), ("addr", "주소", False, 36), ("areaNm", "지역", False, 16),
    ("regionNm", "권역", False, 10), ("smasrNm", "매니저 성함", False, 12), ("smasrHp", "매니저 전화번호", False, 16),
    ("rmk", "비고", False, 30),
]
INVT_AUTO = ["brdNm", "shopFormNm", "shopNm", "prevSaleAmt", "currSaleAmt", "lastInvtDt", "prevInvtType", "prevInvtResult",
             "stockQty", "stockBaseDt", "shopRankNm", "shopTel", "moBrdCd"]


def invt_template() -> bytes:
    auto = ", ".join(invt_plan.LABELS[k] for k in INVT_AUTO if k in invt_plan.LABELS)
    guide = [
        ("필수", "매장코드 (빨간 머리글). 나머지는 비워도 됩니다."),
        ("자동 입력", f"매장코드로 화면에서 매장을 고를 때처럼 자동으로 채웁니다: {auto}. "
                     "자동 값 대신 직접 넣으려면 그 항목 이름으로 열을 추가해 적으세요 (예: '관리등급' 열)."),
        ("실사예정일", "날짜(2026-10-15) 또는 20261015. 비우거나 '미정' 이면 미정."),
        ("정산 팀구분", " / ".join(invt_plan.STLM_TEAMS)),
        ("연2회 실사 매장", "Y 또는 비움"),
        ("지역", " / ".join(invt_plan.AREAS) + " — 권역을 비우면 지역에 맞춰 채웁니다."),
        ("권역", " / ".join(invt_plan.REGIONS)),
        ("비고", f"{invt_plan.RMK_MAX_BYTES}바이트 이내"),
        ("올리기", f"한 번에 {MAX_ROWS['invt']}행까지. 미리보기에서 오류 행을 확인한 뒤 [정상 n건 저장] 을 누르면 등록됩니다 (오류 행은 빠짐)."),
        ("주의", "같은 매장에 이미 등록된 계획이 있으면 '주의' 로 표시만 하고 새 계획으로 추가합니다."),
    ]
    return _template("실사계획 업로드", INVT_COLS, guide,
                     {"stlmTeam": invt_plan.STLM_TEAMS, "twiceYearYn": ["Y", "N"], "regionNm": invt_plan.REGIONS})


def _invt_headers() -> dict[str, str]:
    h = {label: key for key, label in invt_plan.LABELS.items()}
    h.update({label: key for key, label, _, _ in INVT_COLS})
    return h


def _invt_row(no: int, given: dict, shops: dict, existing: dict, seen: dict) -> dict:
    """한 행 검증 + 자동 입력 값. values 는 저장할 값(API 키)."""
    msgs: list[str] = []
    status = "ok"
    sid = str(given.get("shopId") or "").strip().upper()
    vals = {k: v for k, v in given.items() if v is not None}
    if str(vals.get("invtPlanDt") or "").strip() == "미정":
        vals.pop("invtPlanDt")
    if vals.get("twiceYearYn") is not None:
        vals["twiceYearYn"] = "Y" if str(vals["twiceYearYn"]).strip().upper() in ("Y", "1", "예", "O") else "N"
    if vals.get("areaNm") and not vals.get("regionNm"):
        vals["regionNm"] = invt_plan.AREA_REGION.get(vals["areaNm"])
    for k in ("stlmTeam",):   # '1' → '1팀'
        if vals.get(k) is not None and re.fullmatch(r"\d", str(vals[k])):
            vals[k] = f"{vals[k]}팀"
    if not sid:
        return {"row": no, "status": "error", "messages": ["매장코드가 없습니다."], "values": vals}
    vals["shopId"] = sid
    if sid not in shops:
        return {"row": no, "status": "error", "messages": [f"매장(T_SHOP)에 없는 매장코드: {sid}"], "values": vals}
    # 자동 입력: 엑셀에 적은 값이 우선
    auto = shops[sid]
    filled = []
    for k in INVT_AUTO:
        if vals.get(k) in (None, "") and auto["values"].get(k) not in (None, ""):
            vals[k] = auto["values"][k]
            filled.append(k)
    try:
        invt_plan._clean(vals, partial=False)
    except HTTPException as ex:
        return {"row": no, "status": "error", "messages": [_msg(ex)], "values": vals, "auto": filled}
    miss = [invt_plan.LABELS[k] for k in auto.get("missing", []) if vals.get(k) in (None, "")]
    if miss:
        msgs.append(f"자동으로 못 채운 항목: {', '.join(miss)}")
    msgs += auto.get("notes", [])
    if existing.get(sid):
        status = "warn"
        msgs.append(f"이미 등록된 실사계획 {existing[sid]}건이 있습니다 (새 계획으로 추가).")
    if sid in seen:
        status = "warn"
        msgs.append(f"{seen[sid]}행과 같은 매장입니다.")
    seen.setdefault(sid, no)
    return {"row": no, "status": status, "messages": msgs, "values": vals, "auto": filled}


def _invt_rows(items: list[tuple[int, dict]]) -> list[dict]:
    ids = sorted({str(v.get("shopId") or "").strip().upper() for _, v in items} - {""})
    real = {r[0] for r in _in_chunks("SELECT SHOP_ID FROM T_SHOP WHERE SHOP_ID IN ({ids})", ids)} if ids else set()
    existing = {r[0]: int(r[1]) for r in _in_chunks(
        f"SELECT SHOP_ID, COUNT(*) FROM {invt_plan.TABLE} WHERE DEL_DAY IS NULL AND SHOP_ID IN ({{ids}}) GROUP BY SHOP_ID", sorted(real))} if real else {}
    shops: dict[str, dict] = {}

    def detail(sid: str):
        try:
            return sid, invt_plan.shop_detail(sid)
        except Exception as ex:  # noqa: BLE001 - 자동 입력이 실패해도 직접 적은 값으로 등록할 수 있게
            return sid, {"values": {}, "missing": [], "notes": [f"매장 정보 자동 입력 실패: {_msg(ex)}"]}

    with ThreadPoolExecutor(max_workers=4) as pool:
        for sid, d in pool.map(detail, sorted(real)):
            shops[sid] = d
    seen: dict[str, int] = {}
    return [_invt_row(no, v, shops, existing, seen) for no, v in items]


def _display(rows: list[dict]) -> list[dict]:
    for r in rows:
        v = r["values"]
        r["display"] = {k: v.get(k) for k in ("shopId", "shopNm", "brdNm", "shopFormNm", "invtPlanDt", "invtPlanNote", "expectAmt",
                                                "stlmTeam", "shopRankNm", "lastInvtDt")}
    return rows


def invt_preview(body: dict) -> dict:
    items = _read(_decode(body), _invt_headers(), ["shopId"], MAX_ROWS["invt"])
    rows = _display(_invt_rows(items))
    return {"rows": rows, "summary": _summary(rows), "labels": invt_plan.LABELS}


def invt_apply(user: dict, body: dict) -> dict:
    """미리보기에서 받은 행을 다시 검증하고 정상 · 주의 행만 한 트랜잭션으로 등록"""
    raw = (body or {}).get("rows")
    if not isinstance(raw, list) or not raw:
        _bad("저장할 행이 없습니다.")
    if len(raw) > MAX_ROWS["invt"]:
        _bad(f"한 번에 {MAX_ROWS['invt']}행까지 저장할 수 있습니다.")
    items = [(int(r.get("row") or 0), {k: v for k, v in (r.get("values") or {}).items() if k in invt_plan.FIELDS}) for r in raw]
    rows = _invt_rows(items)
    good = [r for r in rows if r["status"] != "error"]
    if not good:
        _bad("저장할 수 있는 행이 없습니다.")
    cols = [c for c, _, _ in invt_plan.FIELDS.values()]
    now = invt_plan._now14()
    params = []
    for r in good:
        vals = invt_plan._clean(r["values"], partial=False)
        vals.setdefault("TWICE_YEAR_YN", "N")
        params.append({**{c: vals.get(c) for c in cols}, "INS_DAY": now, "INS_USERID": user["id"]})
    all_cols = cols + ["INS_DAY", "INS_USERID"]
    sql = (f"INSERT INTO {invt_plan.TABLE} (PLAN_ID, {', '.join(all_cols)}) "
           f"VALUES ({invt_plan.SEQ}.NEXTVAL, {', '.join(':' + c for c in all_cols)})")
    with db.timed(sql, params[0]), db.get_pool().acquire() as conn, conn.cursor() as cur:
        cur.executemany(sql, params)
        conn.commit()
    errors = [{"row": r["row"], "messages": r["messages"]} for r in rows if r["status"] == "error"]
    return {"saved": len(good), "skipped": len(errors), "errors": errors}


# ----------------------------------------------------------------------------
# 판매처 매장 연결
# ----------------------------------------------------------------------------
MALL_COLS = [
    ("mallNm", "사이트명", True, 24), ("sellNo", "판매자번호", False, 16), ("brdCd", "브랜드", False, 12),
    ("shopId", "매장코드", False, 12), ("useYn", "사용", False, 8), ("rmk", "비고", False, 30),
]
BRAND_IN = {**{v: k for k, v in BRAND_CODES.items()}, "모든 브랜드": "*", "전체": "*", "공통": "*"}


def mall_template() -> bytes:
    brands = ", ".join(f"{k}({v})" for k, v in BRAND_CODES.items())
    guide = [
        ("필수", "사이트명 (빨간 머리글). 화면의 '사이트' 와 똑같이 적으세요 (수집 데이터의 MALL_NM)."),
        ("판매자번호", "네이버페이 판매자번호. 없으면 비움 (- 로 저장)."),
        ("브랜드", f"{brands}, 비우거나 * = 모든 브랜드 공통 (그 사이트·판매자에 브랜드 행이 없을 때 쓰임)."),
        ("매장코드", "T_SHOP 의 매장코드. 비우면 그 연결을 해제(삭제)합니다."),
        ("사용", "Y 또는 N. 비우면 기존 연결은 그대로, 새 연결은 Y."),
        ("비고", "비우면 기존 비고를 그대로 둡니다."),
        ("올리기", f"한 번에 {MAX_ROWS['mall']}행까지. 미리보기에서 신규 · 변경 · 해제 · 같음을 확인한 뒤 [정상 n건 저장]. 변경 이력에 남습니다."),
    ]
    return _template("판매처 매장 연결", MALL_COLS, guide, {"brdCd": [*BRAND_CODES, "*"], "useYn": ["Y", "N"]})


def _mall_rows(items: list[tuple[int, dict]]) -> list[dict]:
    clean: list[tuple[int, tuple | None, list[str], dict]] = []
    for no, v in items:
        v = {**v, "_keepUse": v.get("useYn") in (None, ""), "_keepRmk": v.get("rmk") in (None, "")}
        brd = str(v.get("brdCd") or "*").strip()
        v = {**v, "brdCd": BRAND_IN.get(brd, brd.upper()), "sellNo": None if v.get("sellNo") is None else str(v["sellNo"]).strip(),
             "useYn": "N" if str(v.get("useYn") or "Y").strip().upper() in ("N", "아니오", "X") else "Y"}
        try:
            item = mall_shop.clean_item(v)
            clean.append((no, item, [], v))
        except HTTPException as ex:
            clean.append((no, None, [_msg(ex)], v))
    ids = sorted({c[1][3] for c in clean if c[1] and c[1][3]})
    names = mall_shop._shop_names(ids) if ids else {}
    before = mall_shop._maps() if mall_shop.table_ready() else {}
    seen: dict[tuple, int] = {}
    rows = []
    for no, item, errs, v in clean:
        r = {"row": no, "values": v, "messages": list(errs), "status": "error" if errs else "ok", "change": None}
        if item:
            mall, sell, brd, shop, use, rmk = item
            key = (mall, sell, brd)
            b = before.get(key)
            if b and v.get("_keepUse"):     # 비운 칸은 기존 값 유지 (비고를 비웠다고 지우지 않음)
                use = b["USE_YN"]
            if b and v.get("_keepRmk"):
                rmk = b["RMK"] or None
            r["values"] = {"mallNm": mall, "sellNo": sell, "brdCd": brd, "shopId": shop, "useYn": use, "rmk": rmk}
            r["values"]["shopNm"] = names.get(shop) if shop else None
            if shop and shop not in names:
                r["status"], r["messages"] = "error", [f"매장(T_SHOP)에 없는 매장코드: {shop}"]
            elif key in seen:
                r["status"], r["messages"] = "error", [f"{seen[key]}행과 사이트 · 판매자번호 · 브랜드가 같습니다."]
            else:
                seen[key] = no
                if not shop:
                    r["change"] = "해제" if b else "없음"
                    if not b:
                        r["status"], r["messages"] = "warn", ["매장코드가 비어 있는데 기존 연결도 없어 저장할 것이 없습니다."]
                elif not b:
                    r["change"] = "신규"
                elif (b["SHOP_ID"], b["USE_YN"], b["RMK"] or None) == (shop, use, rmk):
                    r["change"] = "같음"
                else:
                    r["change"] = "변경"
                    diff = [f"매장 {b['SHOP_ID']} → {shop}" if b["SHOP_ID"] != shop else "",
                            f"사용 {b['USE_YN']} → {use}" if b["USE_YN"] != use else "",
                            "비고 변경" if (b["RMK"] or None) != rmk else ""]
                    r["messages"].append(" · ".join(d for d in diff if d))
            r["values"]["brand"] = "모든 브랜드" if brd == mall_shop.ALL_BRANDS else BRAND_CODES.get(brd, brd)
        rows.append(r)
    return rows


def mall_preview(body: dict) -> dict:
    headers = {label: key for key, label, _, _ in MALL_COLS} | {"사이트": "mallNm", "매장": "shopId", "사용여부": "useYn"}
    rows = _mall_rows(_read(_decode(body), headers, ["mallNm"], MAX_ROWS["mall"]))
    return {"rows": rows, "summary": {**_summary(rows), "changes": {k: sum(r["change"] == k for r in rows)
                                                                      for k in ("신규", "변경", "해제", "같음")}}}


def mall_apply(admin: dict, body: dict) -> dict:
    raw = (body or {}).get("rows")
    if not isinstance(raw, list) or not raw:
        _bad("저장할 행이 없습니다.")
    keys = ("mallNm", "sellNo", "brdCd", "shopId", "useYn", "rmk")
    rows = _mall_rows([(int(r.get("row") or 0), {k: (r.get("values") or {}).get(k) for k in keys}) for r in raw])
    good = [r for r in rows if r["status"] == "ok" and r["change"] in ("신규", "변경", "해제")]
    if not good:
        return {"saved": 0, "deleted": 0, "changed": 0, "skipped": len(rows)}
    res = mall_shop.save(admin, [{k: r["values"].get(k) for k in keys} for r in good])
    return {**res, "skipped": len(rows) - len(good)}
