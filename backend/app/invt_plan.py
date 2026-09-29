"""데이터 관리 > 매장 재고 실사계획 (T_SHOP_INVT_PLAN)."""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config, db

OWNER = config.DB_OWNER_SCHEMA
TABLE = "T_SHOP_INVT_PLAN"
SEQ = f"{OWNER}.SQ_SHOP_INVT_PLAN"      # SS10DEV 에 시퀀스 시노님이 없어 소유 스키마로 접근
SALE_TABLE = "T_CLOSE_SALE_BASE"

AREAS = [
    "서울특별시", "부산광역시", "대구광역시", "인천광역시", "광주광역시", "대전광역시", "울산광역시",
    "세종특별자치시", "경기도", "강원특별자치도", "충청북도", "충청남도", "전북특별자치도", "전라남도",
    "경상북도", "경상남도", "제주특별자치도",
]
REGIONS = ["수도권", "영남권", "호남권", "충청권", "기타"]
AREA_REGION = {
    "서울특별시": "수도권", "인천광역시": "수도권", "경기도": "수도권",
    "부산광역시": "영남권", "대구광역시": "영남권", "울산광역시": "영남권", "경상북도": "영남권", "경상남도": "영남권",
    "광주광역시": "호남권", "전북특별자치도": "호남권", "전라남도": "호남권",
    "대전광역시": "충청권", "세종특별자치시": "충청권", "충청북도": "충청권", "충청남도": "충청권",
    "강원특별자치도": "기타", "제주특별자치도": "기타",
}
INVT_TYPES = ["교체", "정기", "오픈", "폐점"]
STLM_TEAMS = ["1팀", "2팀"]
RMK_MAX_BYTES = 200

# 편집 가능한 필드: API 키 → (DB 컬럼, 타입, 최대 바이트)
FIELDS: dict[str, tuple[str, str, int | None]] = {
    "shopId": ("SHOP_ID", "str", 20),
    "moBrdCd": ("MO_BRD_CD", "str", 10),
    "brdNm": ("BRD_NM", "str", 50),
    "shopFormNm": ("SHOP_FORM_NM", "str", 100),
    "shopNm": ("SHOP_NM", "str", 200),
    "prevSaleAmt": ("PREV_SALE_AMT", "int", None),
    "currSaleAmt": ("CURR_SALE_AMT", "int", None),
    "addr": ("ADDR", "str", 300),
    "areaNm": ("AREA_NM", "enum:area", None),
    "regionNm": ("REGION_NM", "enum:region", None),
    "lastInvtDt": ("LAST_INVT_DT", "date", None),
    "prevInvtType": ("PREV_INVT_TYPE", "enum:invt_type", None),
    "prevInvtResult": ("PREV_INVT_RESULT", "num", None),
    "stockQty": ("STOCK_QTY", "int", None),
    "stockBaseDt": ("STOCK_BASE_DT", "date", None),
    "invtPlanNote": ("INVT_PLAN_NOTE", "str", 1000),
    "baseFee": ("BASE_FEE", "int", None),
    "expectAmt": ("EXPECT_AMT", "int", None),
    "invtPlanDt": ("INVT_PLAN_DT", "date", None),
    "rmk": ("RMK", "rmk", None),
    "twiceYearYn": ("TWICE_YEAR_YN", "yn", None),
    "shopRankNm": ("SHOP_RANK_NM", "str", 100),
    "stlmTeam": ("STLM_TEAM", "enum:team", None),
    "smasrNm": ("SMASR_NM", "str", 100),
    "smasrHp": ("SMASR_HP", "str", 30),
    "shopTel": ("SHOP_TEL", "str", 30),
}
ENUMS = {"area": AREAS, "region": REGIONS, "invt_type": INVT_TYPES, "team": STLM_TEAMS}
LABELS = {
    "shopId": "매장코드", "brdNm": "브랜드", "shopFormNm": "유통", "shopNm": "매장명", "prevSaleAmt": "전년 매출",
    "currSaleAmt": "당년 매출", "addr": "주소", "areaNm": "지역", "regionNm": "권역", "lastInvtDt": "최종실사일",
    "prevInvtType": "전실사유형", "prevInvtResult": "전실사결과", "stockQty": "재고 수량", "stockBaseDt": "재고 기준일",
    "invtPlanNote": "실사예정", "baseFee": "기본료", "expectAmt": "실사예상액", "invtPlanDt": "실사예정일",
    "rmk": "비고", "twiceYearYn": "연2회 실사 매장", "shopRankNm": "관리등급", "stlmTeam": "정산 팀구분",
    "smasrNm": "매니저 성함", "smasrHp": "매니저 전화번호", "shopTel": "매장번호", "moBrdCd": "브랜드코드",
}


def _bad(msg: str):
    raise HTTPException(status_code=400, detail={"message": msg, "code": "BAD_REQUEST"})


def _now14() -> str:
    return datetime.now().strftime("%Y%m%d%H%M%S")


def _yyyymmdd(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y%m%d")
    s = re.sub(r"[^0-9]", "", str(v))
    return s[:8] if len(s) >= 8 else None


def _num(v):
    if v is None:
        return None
    f = float(v)
    return int(f) if f.is_integer() else round(f, 2)


# ----------------------------------------------------------------------------
# DB 문자셋 (비고 바이트 수 계산)
# ----------------------------------------------------------------------------
_charset: str | None = None


def _db_encoding() -> str:
    global _charset
    if _charset is None:
        try:
            _charset = db.query("SELECT value FROM nls_database_parameters WHERE parameter='NLS_CHARACTERSET'")[1][0][0]
        except Exception:
            _charset = "AL32UTF8"
    return "cp949" if _charset in ("KO16MSWIN949", "KO16KSC5601") else "utf-8"


def byte_len(s: str) -> int:
    return len(s.encode(_db_encoding(), errors="replace"))


# ----------------------------------------------------------------------------
# 목록
# ----------------------------------------------------------------------------
SELECT_COLS = """
    PLAN_ID, SHOP_ID, MO_BRD_CD, BRD_NM, SHOP_FORM_NM, SHOP_NM, PREV_SALE_AMT, CURR_SALE_AMT, ADDR, AREA_NM,
    REGION_NM, LAST_INVT_DT, PREV_INVT_TYPE, PREV_INVT_RESULT, STOCK_QTY, STOCK_BASE_DT, INVT_PLAN_NOTE, BASE_FEE,
    EXPECT_AMT, INVT_PLAN_DT, RMK, TWICE_YEAR_YN, SHOP_RANK_NM, STLM_TEAM, SMASR_NM, SMASR_HP, SHOP_TEL,
    INS_DAY, INS_USERID, UPT_DAY, UPT_USERID,
    CASE WHEN REGEXP_LIKE(LAST_INVT_DT, '^[0-9]{8}$')
         THEN TRUNC(SYSDATE) - TO_DATE(LAST_INVT_DT, 'YYYYMMDD') END AS ELAPSED_DAYS
"""
COL_TO_KEY = {v[0]: k for k, v in FIELDS.items()}


def _to_api(r: dict) -> dict:
    out: dict[str, Any] = {"planId": int(r["PLAN_ID"])}
    for col, key in COL_TO_KEY.items():
        v = r.get(col)
        out[key] = _num(v) if isinstance(v, (int, float)) else v
    prev, curr = out.get("prevSaleAmt"), out.get("currSaleAmt")
    out["prevSaleMil"] = round(prev / 1_000_000) if prev is not None else None
    out["currSaleMil"] = round(curr / 1_000_000) if curr is not None else None
    out["saleRate"] = round((curr / prev - 1) * 100, 1) if prev and curr is not None else None
    out["elapsedDays"] = _num(r.get("ELAPSED_DAYS"))
    out["insDay"], out["insUserId"] = r.get("INS_DAY"), r.get("INS_USERID")
    out["uptDay"], out["uptUserId"] = r.get("UPT_DAY"), r.get("UPT_USERID")
    return out


def list_plans() -> list[dict]:
    rows = db.query_dicts(
        f"SELECT {SELECT_COLS} FROM {TABLE} WHERE DEL_DAY IS NULL ORDER BY INVT_PLAN_DT NULLS LAST, SHOP_ID, PLAN_ID"
    )
    return [_to_api(r) for r in rows]


def get_plan(plan_id: int) -> dict:
    rows = db.query_dicts(f"SELECT {SELECT_COLS} FROM {TABLE} WHERE PLAN_ID = :id AND DEL_DAY IS NULL", {"id": plan_id})
    if not rows:
        raise HTTPException(404, {"message": "실사계획을 찾을 수 없습니다.", "code": "NOT_FOUND"})
    return _to_api(rows[0])


# ----------------------------------------------------------------------------
# 저장
# ----------------------------------------------------------------------------
def _clean(body: dict, partial: bool) -> dict[str, Any]:
    vals: dict[str, Any] = {}
    for key, (col, typ, maxlen) in FIELDS.items():
        if key not in body:
            continue
        v = body[key]
        label = LABELS.get(key, key)
        if isinstance(v, str):
            v = v.strip()
        if v == "" or v is None:
            v = None
        elif typ == "str":
            v = str(v)
            if maxlen and byte_len(v) > maxlen:
                _bad(f"{label}은(는) {maxlen}바이트 이내로 입력하세요.")
        elif typ in ("int", "num"):
            try:
                f = float(str(v).replace(",", ""))
            except ValueError:
                _bad(f"{label}은(는) 숫자로 입력하세요.")
            v = int(round(f)) if typ == "int" else round(f, 2)
            if abs(v) >= 10 ** 12:
                _bad(f"{label} 값이 너무 큽니다.")
        elif typ == "date":
            d = _yyyymmdd(v)
            try:
                datetime.strptime(d or "", "%Y%m%d")
            except ValueError:
                _bad(f"{label}은(는) 올바른 날짜(YYYYMMDD)가 아닙니다.")
            v = d
        elif typ.startswith("enum:"):
            if v not in ENUMS[typ[5:]]:
                _bad(f"{label} 값이 올바르지 않습니다: {v}")
        elif typ == "yn":
            v = "Y" if v in (True, "Y", "y", 1, "1") else "N"
        elif typ == "rmk":
            v = str(v)
            if byte_len(v) > RMK_MAX_BYTES:
                _bad(f"비고는 {RMK_MAX_BYTES}바이트 이내로 입력하세요. (현재 {byte_len(v)}바이트)")
        vals[col] = v
    if "TWICE_YEAR_YN" in vals and vals["TWICE_YEAR_YN"] is None:
        vals["TWICE_YEAR_YN"] = "N"
    if not partial and not vals.get("SHOP_ID"):
        _bad("매장코드를 선택하세요.")
    if "SHOP_ID" in vals and vals["SHOP_ID"] is None:
        _bad("매장코드는 비울 수 없습니다.")
    return vals


def create_plan(user_id: str, body: dict) -> dict:
    vals = _clean(body, partial=False)
    vals.setdefault("TWICE_YEAR_YN", "N")
    new_id = int(db.query(f"SELECT {SEQ}.NEXTVAL FROM DUAL")[1][0][0])
    vals.update({"PLAN_ID": new_id, "INS_DAY": _now14(), "INS_USERID": user_id})
    cols = list(vals)
    db.execute(f"INSERT INTO {TABLE} ({', '.join(cols)}) VALUES ({', '.join(':' + c for c in cols)})", vals)
    return get_plan(new_id)


def update_plan(user_id: str, plan_id: int, body: dict) -> dict:
    vals = _clean(body, partial=True)
    if vals:
        vals.update({"UPT_DAY": _now14(), "UPT_USERID": user_id})
        sets = ", ".join(f"{c} = :{c}" for c in vals)
        n = db.execute(f"UPDATE {TABLE} SET {sets} WHERE PLAN_ID = :PLAN_ID AND DEL_DAY IS NULL",
                       {**vals, "PLAN_ID": plan_id})
        if n == 0:
            raise HTTPException(404, {"message": "실사계획을 찾을 수 없습니다.", "code": "NOT_FOUND"})
    return get_plan(plan_id)


def delete_plans(user_id: str, plan_ids: list[int]) -> int:
    if not plan_ids:
        return 0
    binds = {f"id{i}": int(v) for i, v in enumerate(plan_ids[:500])}
    return db.execute(
        f"UPDATE {TABLE} SET DEL_DAY = :d, DEL_USERID = :u "
        f"WHERE DEL_DAY IS NULL AND PLAN_ID IN ({', '.join(':' + k for k in binds)})",
        {"d": _now14(), "u": user_id, **binds},
    )


def active_count_for_shop(shop_id: str) -> int:
    return int(db.query(f"SELECT COUNT(*) FROM {TABLE} WHERE SHOP_ID = :s AND DEL_DAY IS NULL", {"s": shop_id})[1][0][0])


# ----------------------------------------------------------------------------
# 매장 / 매니저 팝업
# ----------------------------------------------------------------------------
def search_shops(q: str) -> list[dict]:
    q = (q or "").strip()
    if not q:
        return []
    rows = db.query_dicts(
        """
        SELECT * FROM (
            SELECT SHOP_ID, SHOP_NM,
                   F_GET_CD_NM(SHOP_FORM, 'ko')  AS SHOP_FORM_NM,
                   F_GET_CD_NM(SHOP_FORM2, 'ko') AS SHOP_FORM2_NM
              FROM T_SHOP
             WHERE UPPER(SHOP_ID) LIKE UPPER(:q) OR SHOP_NM LIKE :q
             ORDER BY CASE WHEN UPPER(SHOP_ID) = UPPER(:exact) THEN 0 ELSE 1 END, SHOP_ID
        ) WHERE ROWNUM <= 100
        """,
        {"q": f"%{q}%", "exact": q},
    )
    return [{"shopId": r["SHOP_ID"], "shopNm": r["SHOP_NM"], "shopFormNm": r["SHOP_FORM_NM"],
             "shopForm2Nm": r["SHOP_FORM2_NM"]} for r in rows]


def _safe(fn):
    try:
        return fn(), None
    except Exception as ex:  # 한 항목 조회 실패가 전체 자동 세팅을 막지 않도록
        return None, str(ex)


def invt_rank(shop_id: str, invt_dt: str | None) -> str | None:
    """관리등급: 최종실사일 기준 T_SHOP_INVT_RANK.FINAL_RANK."""
    if not invt_dt:
        return None
    rows = db.query(
        "SELECT FINAL_RANK FROM T_SHOP_INVT_RANK WHERE SHOP_ID = :id AND INVT_DT = :dt",
        {"id": shop_id, "dt": invt_dt},
    )[1]
    return str(rows[0][0]) if rows and rows[0][0] is not None else None


def shop_detail(shop_id: str) -> dict:
    """매장 선택 시 자동 세팅 값. 조회되지 않은 항목은 missing 에 담아 수기 입력을 안내한다."""
    p = {"id": shop_id}
    errors: dict[str, str] = {}

    base, err = _safe(lambda: db.query_dicts(
        """
        SELECT S.SHOP_ID, S.MO_BRD_CD,
               DECODE(S.MO_BRD_CD, 'S', '쉬즈미스', 'T', '리스트', 'A', '시스티나') AS BRD_NM,
               F_GET_CD_NM(S.SHOP_FORM, 'ko') AS SHOP_FORM_NM,
               S.SHOP_NM,
               S.SHOP_TEL_NO1, S.SHOP_TEL_NO2, S.SHOP_TEL_NO3
          FROM T_SHOP S
         WHERE S.SHOP_ID = :id
        """, p))
    if err:
        errors["base"] = err
    if not base and not err:
        raise HTTPException(404, {"message": f"매장코드 {shop_id} 를 찾을 수 없습니다.", "code": "NOT_FOUND"})
    b = base[0] if base else {}

    # 재고 수량: 당월 재고 (재고 행이 없어도 매장 정보는 나오도록 별도 조회)
    stock, err = _safe(lambda: db.query(
        """SELECT SUM(STOCK_QTY) FROM T_SHOP_STOCK
            WHERE SHOP_ID = :id AND MAKE_YYMM = TO_CHAR(SYSDATE, 'YYYYMM') AND STOCK_QTY <> 0""", p)[1][0][0])
    if err:
        errors["stock"] = err

    # 최종실사일 · 전실사유형 · 전실사결과: 최종 실사일을 먼저 구한 뒤(WITH), 그 실사 1건 기준으로 계산
    # (원 쿼리처럼 전체 기간을 한 번에 집계하면 과거 실사 결과가 모두 합산되고 유형도 최신 건이 아님)
    last, err = _safe(lambda: db.query_dicts(
        """
        WITH LAST_INVT AS (
            SELECT MAX(INV.INVT_DT) AS INVT_DT
              FROM T_SHOP_INVT_STLM_TOT TOT
                 , T_SHOP_INVT          INV
                 , T_SHOP_INVT_PRE_INFO PRE
             WHERE INV.SHOP_ID = TOT.SHOP_ID
               AND INV.SHOP_ID = PRE.SHOP_ID
               AND INV.INVT_DT = TOT.INVT_DT
               AND INV.INVT_DT = PRE.INVT_EXEC_DT
               AND INV.INVT_COST_CLSBY = 'C70210'
               AND TOT.SHOP_ID = :id
        )
        SELECT L.INVT_DT AS LAST_INVT_DT
             , (SELECT F_CD_NM(MAX(PRE.INVT_EXEC_CLSBY))
                  FROM T_SHOP_INVT_PRE_INFO PRE
                 WHERE PRE.SHOP_ID = :id
                   AND PRE.INVT_EXEC_DT = L.INVT_DT) AS PREV_INVT_TYPE
             , (SELECT SUM(TOT.INVT_DFNT_QTY * TOT.SUPP_RPICE) - SUM(TOT.SYSTEM_STOCK_QTY * TOT.SUPP_RPICE)
                  FROM T_SHOP_INVT_STLM_TOT TOT
                 WHERE TOT.SHOP_ID = :id
                   AND TOT.INVT_DT = L.INVT_DT) AS PREV_INVT_RESULT
          FROM LAST_INVT L
        """, p))
    if err:
        errors["lastInvt"] = err
    li = last[0] if last else {}
    last_dt_s = _yyyymmdd(li.get("LAST_INVT_DT"))
    prev_type = li.get("PREV_INVT_TYPE") if last_dt_s else None
    prev_result = _num(li.get("PREV_INVT_RESULT")) if last_dt_s else None
    notes: list[str] = []
    if prev_type and prev_type not in INVT_TYPES:
        # 테이블 CHECK 제약(교체/정기/오픈/폐점)에 없는 유형은 저장할 수 없어 비워 두고 안내
        notes.append(f"전실사유형 '{prev_type}'은(는) 선택 목록({'/'.join(INVT_TYPES)})에 없어 비워 두었습니다.")
        prev_type = None

    rank, err = _safe(lambda: invt_rank(shop_id, last_dt_s))
    if err:
        errors["rank"] = err

    sales, err = _safe(lambda: db.query_dicts(
        f"""
        SELECT SUM(CASE WHEN MAKE_YYMM BETWEEN TO_CHAR(SYSDATE, 'YYYY') || '01'
                                          AND TO_CHAR(ADD_MONTHS(SYSDATE, -1), 'YYYYMM')
                        THEN REAL_SALE_AMT END) AS CURR_SALE,
               SUM(CASE WHEN MAKE_YYMM BETWEEN TO_CHAR(ADD_MONTHS(SYSDATE, -12), 'YYYY') || '01'
                                          AND TO_CHAR(ADD_MONTHS(SYSDATE, -13), 'YYYYMM')
                        THEN REAL_SALE_AMT END) AS PREV_SALE
          FROM {SALE_TABLE}
         WHERE SHOP_ID = :id
           AND (MAKE_YYMM BETWEEN TO_CHAR(SYSDATE, 'YYYY') || '01' AND TO_CHAR(ADD_MONTHS(SYSDATE, -1), 'YYYYMM')
                OR MAKE_YYMM BETWEEN TO_CHAR(ADD_MONTHS(SYSDATE, -12), 'YYYY') || '01'
                                 AND TO_CHAR(ADD_MONTHS(SYSDATE, -13), 'YYYYMM'))
        """, p))
    if err:
        errors["sales"] = err
    s = sales[0] if sales else {}

    tel = "-".join(x for x in (str(b.get(k) or "").strip() for k in ("SHOP_TEL_NO1", "SHOP_TEL_NO2", "SHOP_TEL_NO3")) if x)

    values = {
        "shopId": shop_id,
        "moBrdCd": b.get("MO_BRD_CD"),
        "brdNm": b.get("BRD_NM"),
        "shopFormNm": b.get("SHOP_FORM_NM"),
        "shopNm": b.get("SHOP_NM"),
        "shopRankNm": rank,
        "shopTel": tel or None,
        "stockQty": _num(stock),
        "stockBaseDt": date.today().strftime("%Y%m%d") if stock is not None else None,
        "lastInvtDt": last_dt_s,
        "prevInvtType": prev_type,
        "prevInvtResult": prev_result,
        "prevSaleAmt": _num(s.get("PREV_SALE")),
        "currSaleAmt": _num(s.get("CURR_SALE")),
    }
    auto_keys = ["brdNm", "shopFormNm", "shopNm", "prevSaleAmt", "currSaleAmt", "lastInvtDt", "prevInvtType",
                 "prevInvtResult", "stockQty", "shopRankNm", "shopTel"]
    missing = [k for k in auto_keys if values.get(k) in (None, "")]
    return {
        "values": values,
        "missing": missing,
        "missingLabels": [LABELS[k] for k in missing],
        "notes": notes,
        "errors": errors,
        "existingPlans": active_count_for_shop(shop_id),
    }


def shop_managers(shop_id: str) -> list[dict]:
    """매니저 팝업: 최신 등록(SMASR_ID) 순. 종료일 99991231 = 현재 근무."""
    rows = db.query_dicts(
        """
        SELECT SHOP_ID
             , SMASR_NM
             , HP_NO1 || '-' || HP_NO2 || '-' || HP_NO3 AS SMASR_HP
             , OPEN_DT
             , CLOSE_DT
          FROM T_SHOP_SMAS
         WHERE SHOP_ID = :id
         ORDER BY SMASR_ID DESC
        """,
        {"id": shop_id},
    )
    out = []
    for r in rows:
        hp = (r.get("SMASR_HP") or "").strip("- ")  # 번호가 비어 있으면 '--' 가 되므로 정리
        close_dt = r.get("CLOSE_DT")
        out.append({
            "shopId": r["SHOP_ID"],
            "smasrNm": r["SMASR_NM"],
            "smasrHp": hp or None,
            "openDt": r.get("OPEN_DT"),
            "closeDt": close_dt,
            "current": close_dt is None or close_dt >= "99991231" or close_dt >= date.today().strftime("%Y%m%d"),
        })
    return out


# ----------------------------------------------------------------------------
# 엑셀 (화면과 같은 2단 헤더)
# ----------------------------------------------------------------------------
# (그룹명 또는 None, 라벨, 키, 너비, 형식)
EXCEL_COLS: list[tuple[str | None, str, str, int, str]] = [
    (None, "매장코드", "shopId", 11, "text"),
    (None, "브랜드", "brdNm", 10, "text"),
    (None, "유통", "shopFormNm", 12, "text"),
    (None, "매장명", "shopNm", 22, "text"),
    ("평균매출(백만원)", "전년", "prevSaleMil", 9, "int"),
    ("평균매출(백만원)", "당년", "currSaleMil", 9, "int"),
    ("평균매출(백만원)", "증감율(%)", "saleRate", 10, "pct"),
    (None, "주소", "addr", 36, "text"),
    (None, "지역", "areaNm", 13, "text"),
    (None, "권역", "regionNm", 9, "text"),
    (None, "최종실사일", "lastInvtDt", 12, "date"),
    (None, "전실사유형", "prevInvtType", 10, "text"),
    (None, "전실사결과", "prevInvtResult", 11, "num"),
    (None, "경과일", "elapsedDays", 8, "int"),
    ("재고 수량", "당일기준", "stockQty", 11, "int"),
    (None, "실사예정", "invtPlanNote", 24, "text"),
    ("업체 예상 비용", "기본료", "baseFee", 12, "int"),
    ("업체 예상 비용", "실사예상액", "expectAmt", 12, "int"),
    (None, "실사예정일", "invtPlanDt", 12, "date"),
    (None, "비고", "rmk", 30, "text"),
    (None, "연2회 실사 매장", "twiceYearYn", 10, "yn"),
    (None, "관리등급", "shopRankNm", 10, "text"),
    (None, "정산 팀구분", "stlmTeam", 10, "text"),
    ("매니저", "성함", "smasrNm", 10, "text"),
    ("매니저", "전화번호", "smasrHp", 15, "text"),
    ("매니저", "매장번호", "shopTel", 15, "text"),
]


def export_xlsx(plans: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "매장 재고 실사계획"
    head_fill = PatternFill("solid", fgColor="1F1F7A")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="8A8FB3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    col, i = 1, 0
    while i < len(EXCEL_COLS):
        group = EXCEL_COLS[i][0]
        span = 1
        while group and i + span < len(EXCEL_COLS) and EXCEL_COLS[i + span][0] == group:
            span += 1
        if group:
            ws.cell(1, col, group)
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + span - 1)
            for j in range(span):
                ws.cell(2, col + j, EXCEL_COLS[i + j][1])
        else:
            ws.cell(1, col, EXCEL_COLS[i][1])
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        for j in range(span):
            ws.column_dimensions[get_column_letter(col + j)].width = EXCEL_COLS[i + j][3]
        col += span
        i += span
    for r in (1, 2):
        ws.row_dimensions[r].height = 20
        for c in range(1, len(EXCEL_COLS) + 1):
            cell = ws.cell(r, c)
            cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, center, border
    ws.freeze_panes = "E3"

    for ri, p in enumerate(plans, start=3):
        for ci, (_, _, key, _, fmt) in enumerate(EXCEL_COLS, start=1):
            v = p.get(key)
            if key == "invtPlanDt" and not v:
                v = "미정"
            elif fmt == "date" and v:
                v = f"{v[:4]}-{v[4:6]}-{v[6:8]}"
            elif fmt == "yn":
                v = "Y" if v == "Y" else ""
            cell = ws.cell(ri, ci, v)
            cell.border = border
            if fmt in ("int", "num") and isinstance(v, (int, float)):
                cell.number_format = "#,##0" if fmt == "int" else "#,##0.##"
            elif fmt == "pct" and isinstance(v, (int, float)):
                cell.number_format = "0.0"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def options() -> dict:
    return {
        "areas": AREAS,
        "regions": REGIONS,
        "areaRegion": AREA_REGION,
        "invtTypes": INVT_TYPES,
        "stlmTeams": STLM_TEAMS,
        "rmkMaxBytes": RMK_MAX_BYTES,
        "encoding": _db_encoding(),
    }
