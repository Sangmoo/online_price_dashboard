"""판매 현황 보고용 엑셀: 화면과 같은 조건·같은 숫자를 시트별로 정리한다.

시트: 요약(차트 포함) · 월별 추이 · 브랜드별 · 팀별 · 매장 순위(매출 상위/성장/하락/목표 미달) · 전체 매장
차트는 엑셀 차트라 데이터 시트 값을 그대로 참조한다 (엑셀에서 바로 편집·복사 가능).
금액은 원 단위(#,##0), 비율은 % (0.0). 인쇄·보고에 바로 쓰도록 제목·조건·작성 시각을 위에 적는다.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import sale_dashboard as sd

TITLE_FONT = Font(name="맑은 고딕", size=14, bold=True)
NOTE_FONT = Font(name="맑은 고딕", size=9, color="FF6B7280")
HEAD_FONT = Font(name="맑은 고딕", size=10, bold=True)
BODY_FONT = Font(name="맑은 고딕", size=10)
SECTION_FONT = Font(name="맑은 고딕", size=11, bold=True, color="FF3730A3")
HEAD_FILL = PatternFill("solid", fgColor="FFE8ECF7")
THIN = Side(style="thin", color="FFB4BCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FMT = {"amt": "#,##0", "pct": '0.0"%"', "pp": '+0.0"%p";-0.0"%p";0.0"%p"', "chg": '+0.0"%";-0.0"%";0.0"%"', "int": "#,##0",
       "text": "@"}


def _header(ws, d: dict, title: str) -> int:
    """제목 + 조건 줄. 다음에 쓸 행 번호를 돌려준다."""
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    cond = (f"기간 {d['period']['label']} · 비교 {d['base']['kindLabel']} {d['base']['label']} · 브랜드 {d['brand'] or '전체'} · "
            f"마감 매출(실판금액) 기준 · 작성 {datetime.now():%Y-%m-%d %H:%M}")
    ws["A2"] = cond
    ws["A2"].font = NOTE_FONT
    return 4


def _table(ws, row: int, cols: list[tuple[str, str, str, float]], rows: list[dict], section: str | None = None) -> int:
    """cols: (키, 제목, 서식, 너비). 표를 쓰고 다음 빈 행 번호(한 줄 띄움)를 돌려준다."""
    if section:
        ws.cell(row=row, column=1, value=section).font = SECTION_FONT
        row += 1
    for i, (_, label, _, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)
    for r in rows:
        row += 1
        for i, (key, _, fmt, _) in enumerate(cols, start=1):
            v = r.get(key)
            c = ws.cell(row=row, column=i, value=v)
            c.font, c.border = BODY_FONT, BORDER
            if v is not None and fmt != "text":
                c.number_format = FMT[fmt]
    if not rows:
        row += 1
        ws.cell(row=row, column=1, value="(해당 없음)").font = NOTE_FONT
    return row + 2


def _kpi_rows(d: dict) -> list[dict]:
    k, bl, el = d["kpi"], f"{d['base']['kindLabel']} {d['base']['label']}", f"{d['extra']['label']} {d['extra']['period']}"
    rows = [
        {"item": "실판금액(원)", "cur": k["amt"], "basis": bl, "base": k["baseAmt"], "chg": k["change"]},
        {"item": "실판금액(원)", "cur": k["amt"], "basis": el, "base": k["extraAmt"], "chg": k["extraChange"]},
        {"item": f"{d['ym'][:4]}년 누계(원)", "cur": k["ytdAmt"], "basis": "전년 같은 기간", "base": k["prevYtdAmt"], "chg": k["ytdYoy"]},
        {"item": "판매 수량", "cur": k["qty"], "basis": bl, "base": k["baseQty"], "chg": k["qtyChange"]},
        {"item": "할인금액(원)", "cur": k["dsct"], "basis": bl, "base": k["baseDsct"], "chg": k["dsctChange"]},
        {"item": "원가 금액(원)", "cur": k["cost"], "basis": "", "base": None, "chg": None},
        {"item": "판매 매장 수", "cur": k["shops"], "basis": bl, "base": k["baseShops"], "chg": None},
        {"item": "매장당 실판금액(원)", "cur": k["avgPerShop"], "basis": "", "base": None, "chg": None},
    ]
    return rows


def build(d: dict) -> bytes:
    """d: sale_dashboard.dashboard(..., full=True) 결과"""
    wb = Workbook()
    k = d["kpi"]
    base_label = f"비교({d['base']['label']})"

    # 1) 요약
    ws = wb.active
    ws.title = "요약"
    row = _header(ws, d, "판매 현황 보고")
    row = _table(ws, row, [("item", "항목", "text", 22), ("cur", d["period"]["label"], "amt", 18), ("basis", "비교 기준", "text", 26),
                           ("base", "비교 값", "amt", 18), ("chg", "증감(%)", "chg", 11)], _kpi_rows(d), "핵심 지표")
    rate_rows = [{"item": "원가율", "cur": k["costRate"], "base": k["baseCostRate"], "chg": k["costRateDiff"]},
                 {"item": "할인율", "cur": k.get("dsctRate"), "base": k.get("baseDsctRate"), "chg": k.get("dsctRateDiff")}]
    row = _table(ws, row, [("item", "항목", "text", 22), ("cur", d["period"]["label"], "pct", 18), ("base", base_label, "pct", 26),
                           ("chg", "증감(%p)", "pp", 11)], rate_rows, "원가율 (원가 금액 ÷ 실판금액) · 할인율 (할인금액 ÷ (실판금액 + 할인금액))")
    if d["hasGoals"]:
        goal_rows = [
            {"item": "목표금액(원)", "v": k["goalAmt"]},
            {"item": "목표 매장 실판금액(원)", "v": k["goalSalesAmt"]},
            {"item": "목표 대비 차이(원)", "v": k["goalGap"]},
            {"item": "목표 매장 수", "v": k["goalShops"]},
            {"item": "목표 없는 매장 수 / 실판금액(원)", "v": k["noGoalShops"], "v2": k["noGoalAmt"]},
        ]
        row = _table(ws, row, [("item", "항목", "text", 22), ("v", "값", "amt", 18), ("v2", "", "amt", 26)], goal_rows,
                     f"목표 대비 · 달성률 {k['achieve'] if k['achieve'] is not None else '-'}%")
        ws.cell(row=row - 1, column=1, value="달성률 = 목표가 있는 매장의 실판금액 ÷ 목표금액. 매장별 목표를 매장의 판매 브랜드로 모음.").font = NOTE_FONT

    # 2) 월별 추이
    ws = wb.create_sheet("월별 추이")
    row = _header(ws, d, f"최근 13개월 실판금액 ({sd.ym_label(d['ym'])}까지)")
    trend = [{**t, "ymLabel": sd.ym_label(t["ym"])} for t in d["trend"]]
    _table(ws, row, [("ymLabel", "판매년월", "text", 12), ("amt", "실판금액", "amt", 18), ("prevAmt", "전년 같은 달", "amt", 18),
                     ("yoy", "전년 대비(%)", "chg", 12), ("costRate", "원가율(%)", "pct", 11), ("dsctRate", "할인율(%)", "pct", 11),
                     ("shops", "판매 매장 수", "int", 12)], trend)

    grp_cols = [("amt", "실판금액", "amt", 18), ("baseAmt", base_label, "amt", 18), ("change", "증감(%)", "chg", 11),
                ("share", "비중(%)", "pct", 10), ("costRate", "원가율(%)", "pct", 11), ("shops", "매장 수", "int", 9),
                ("goalAmt", "목표금액", "amt", 18), ("achieve", "달성률(%)", "pct", 11), ("dsctRate", "할인율(%)", "pct", 10)]
    # 3) 브랜드별
    ws = wb.create_sheet("브랜드별")
    row = _header(ws, d, "브랜드별 실적")
    _table(ws, row, [("brand", "브랜드", "text", 14)] + grp_cols, d["brands"])
    # 4) 팀별
    ws = wb.create_sheet("팀별")
    row = _header(ws, d, "팀별 실적")
    _table(ws, row, [("brand", "브랜드", "text", 12), ("team", "팀", "text", 14)] + [c for c in grp_cols if c[0] != "share"], d["teams"])

    shop_cols = [("rank", "순위", "int", 6), ("shopNm", "매장명", "text", 22), ("shopId", "매장코드", "text", 10), ("brand", "브랜드", "text", 10),
                 ("team", "팀", "text", 12), ("amt", "실판금액", "amt", 16), ("baseAmt", base_label, "amt", 18),
                 ("change", "증감(%)", "chg", 10), ("costRate", "원가율(%)", "pct", 10), ("goalAmt", "목표금액", "amt", 16),
                 ("achieve", "달성률(%)", "pct", 10)]
    ranked = lambda xs: [{**x, "rank": i + 1} for i, x in enumerate(xs)]  # noqa: E731
    # 5) 매장 순위
    ws = wb.create_sheet("매장 순위")
    row = _header(ws, d, "매장 순위")
    row = _table(ws, row, shop_cols, ranked(d["topShops"]), "실판금액 상위 10")
    rule = f"(비교 기간 실판금액 {d['minBaseForGrowth']:,}원 이상 · 폐점·종료 표시 매장 {d['shopCounts']['closedExcluded']}개 제외)"
    row = _table(ws, row, shop_cols, ranked(d["risers"]), f"{d['base']['kindLabel']} 대비 성장 상위 10 {rule}")
    row = _table(ws, row, shop_cols, ranked(d["fallers"]), f"{d['base']['kindLabel']} 대비 하락 상위 10 {rule}")
    if d["hasGoals"]:
        _table(ws, row, shop_cols, ranked(d["laggards"]), "목표 달성률 하위 10 (목표·매출이 있는 영업 매장)")
    _charts(wb, d, base_label)
    # 6) 전체 매장
    ws = wb.create_sheet("전체 매장")
    row = _header(ws, d, f"전체 매장 ({len(d['allShops']):,}개, 실판금액 순)")
    _table(ws, row, shop_cols, ranked(d["allShops"]))
    for s in wb.worksheets:
        s.freeze_panes = None
        s.sheet_view.showGridLines = False

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DATA_ROW = 5  # 데이터 시트의 첫 데이터 행 (1 제목, 2 조건, 4 머리글)
MIL_FMT = '#,##0,,"백만"'


def _axes(chart) -> None:
    # openpyxl 3.1 은 축을 기본으로 숨긴다 → 보이게
    chart.x_axis.delete = False
    chart.y_axis.delete = False


def _charts(wb: Workbook, d: dict, base_label: str) -> None:
    """요약 시트 오른쪽에 월별 추이(막대 당해 + 선 전년)와 브랜드별(당기·비교·목표) 차트"""
    ws = wb["요약"]
    n = len(d["trend"])
    if n:
        src = wb["월별 추이"]
        cats = Reference(src, min_col=1, min_row=DATA_ROW, max_row=DATA_ROW + n - 1)
        bar = BarChart()
        bar.title = "최근 13개월 실판금액 (막대 당해 · 선 전년 같은 달)"
        bar.add_data(Reference(src, min_col=2, min_row=DATA_ROW - 1, max_row=DATA_ROW + n - 1), titles_from_data=True)
        bar.set_categories(cats)
        bar.y_axis.number_format = MIL_FMT
        bar.y_axis.majorGridlines = None
        line = LineChart()
        line.add_data(Reference(src, min_col=3, min_row=DATA_ROW - 1, max_row=DATA_ROW + n - 1), titles_from_data=True)
        line.set_categories(cats)
        bar += line
        _axes(bar)
        bar.height, bar.width = 8, 18
        bar.legend.position = "b"
        ws.add_chart(bar, "G4")
    m = len(d["brands"])
    if m:
        src = wb["브랜드별"]
        chart = BarChart()
        chart.title = f"브랜드별 실판금액 ({d['period']['label']} · {base_label}{' · 목표' if d['hasGoals'] else ''})"
        for col in (2, 3) + ((8,) if d["hasGoals"] else ()):  # 실판금액, 비교, 목표금액
            chart.add_data(Reference(src, min_col=col, min_row=DATA_ROW - 1, max_row=DATA_ROW + m - 1), titles_from_data=True)
        chart.set_categories(Reference(src, min_col=1, min_row=DATA_ROW, max_row=DATA_ROW + m - 1))
        chart.y_axis.number_format = MIL_FMT
        _axes(chart)
        chart.height, chart.width = 8, 18
        chart.legend.position = "b"
        ws.add_chart(chart, "G21")


def filename(d: dict) -> str:
    p = d["period"]
    span = p["from"] if p["from"] == p["to"] else f"{p['from']}-{p['to']}"
    return f"판매현황_{span}_{d['base']['kind']}{'_' + d['brand'] if d['brand'] else ''}.xlsx"
