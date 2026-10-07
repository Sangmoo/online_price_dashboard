"""재고 재배치 추천(매장 간 RT · 창고 → 매장 배분) 공용: 브랜드 · 기간 확인, 매장 기준 정보, 수불제어 판정.

수불제어는 ERP 함수와 같은 규칙을 파이썬으로 계산한다 (함수는 호출당 약 30ms 라 후보 수천 건에 쓰면 몇 분이 걸림).
- controlled_id(...)   = SS10DEV.F_GET_RNDS_CNTR_ID 가 값을 돌려주는지 (F_GET_RNDS_CNTR 의 '13' 판매분 자동보충 · '36' 자동RT 반입)
- controlled_rt_out(...) = SS10DEV.F_GET_RNDS_CNTR_AUTO_RT(…, '32') = 'Y' (자동 RT 반출: 제어 + 자동RT 제외 스타일 + 매장 요청가능여부)
유효 제어(오늘 기준)는 브랜드마다 수백 건이라 5분 캐시로 한 번에 읽는다. 함수와 같은 결과인지는 db/검증 스크립트 대신
단위 테스트(규칙) + 운영 DB 표본 대조로 확인했다 (2026-10-07, 쉬즈미스 · 리스트 · 시스티나 표본 421건 모두 일치).
"""
from __future__ import annotations

import re
import threading
import time
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from fastapi import HTTPException

from . import db
from .shop_info import BRAND_CODES

COMPY_CD = "A01C01"
WEB_MARK = "webStockRt"   # 이 화면에서 ERP 에 등록한 행 표시 (T_INDC_RT.ATTR1 · T_DELV_ASK.ATTR1) — 이 표시가 있는 미확정 행만 삭제할 수 있다
MAX_DAYS = 31                       # 조회 기간 최대 (한 달)
CACHE_TTL = 10 * 60
CTL_TTL = 5 * 60
_D8 = re.compile(r"^\d{8}$")
_lock = threading.Lock()
_cache: dict[tuple, tuple[float, object]] = {}
_key_locks: dict[tuple, threading.Lock] = {}

# F_GET_RNDS_CNTR_ID 의 제어구분별 플래그 (첫 번째 쿼리 = 매장 지정, 두 번째 = 제품 지정). 원본 DECODE 그대로 (34 · 36 은 두 번째에서 INDC_RT_CNTR_YN, 35 는 없음)
_FLAGS1 = {"13": ("AUTO_DVID_CNTR_YN", "DISTRB_DELV_CNTR_YN"), "36": ("AUTO_RT_CNTR_STOR_YN",)}
_FLAGS2 = {"13": ("AUTO_DVID_CNTR_YN", "DISTRB_DELV_CNTR_YN"), "36": ("INDC_RT_CNTR_YN",)}
_FLAG_COLS = ("AUTO_DVID_CNTR_YN", "DISTRB_DELV_CNTR_YN", "AUTO_RT_CNTR_STOR_YN", "INDC_RT_CNTR_YN", "AUTO_RT_CNTR_YN")


def bad(msg: str, status: int = 400):
    raise HTTPException(status_code=status, detail={"message": msg, "code": "BAD_REQUEST" if status == 400 else "FORBIDDEN"})


def cached(key: tuple, ttl: int, fn, force: bool = False):
    """키별 캐시. 같은 키를 동시에 계산하지 않는다 (추천은 수십 초라 두 번 누르면 DB 부하만 두 배)"""
    with _lock:
        hit = _cache.get(key)
        if hit and hit[0] > time.time() and not force:
            return hit[1]
        klock = _key_locks.setdefault(key, threading.Lock())
    with klock:
        with _lock:
            hit = _cache.get(key)
            if hit and hit[0] > time.time() and not force:
                return hit[1]
        val = fn()
        with _lock:
            if len(_cache) > 200:
                _cache.clear()
                _key_locks.clear()
            _cache[key] = (time.time() + ttl, val)
        return val


def cache_time(key: tuple) -> float | None:
    with _lock:
        hit = _cache.get(key)
        return hit[0] if hit else None


def clear_cache() -> None:
    with _lock:
        _cache.clear()
        _key_locks.clear()


def drop_cache(*tags: str) -> None:
    """키 첫 값이 tags 인 캐시만 비운다 (ERP 에 지시 · 의뢰를 등록 · 삭제한 뒤 추천을 다시 계산하게)"""
    with _lock:
        for k in [k for k in _cache if k and k[0] in tags]:
            _cache.pop(k, None)


def xlsx(sheets: list[tuple[str, list[str], list[tuple], list[list]]]) -> bytes:
    """시트 목록 [(이름, 위쪽 안내 줄, [(키, 제목, 너비)], 행)] → 엑셀 바이트"""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F1F7A")
    for title, notes, cols, rows in sheets:
        ws = wb.create_sheet(title[:31])
        for i, line in enumerate(notes, 1):
            ws.cell(i, 1, line).font = Font(bold=i == 1, size=13 if i == 1 else 10)
        hr = len(notes) + 2
        for j, (_, label, width) in enumerate(cols, 1):
            c = ws.cell(hr, j, label)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = width
        for i, r in enumerate(rows, hr + 1):
            for j, (key, _, _) in enumerate(cols, 1):
                v = r.get(key)
                ws.cell(i, j, ("Y" if v else "") if isinstance(v, bool) else v)
        ws.freeze_panes = ws.cell(hr + 1, 1)
        if rows:
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{hr + len(rows)}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def default_seasons(now: date | None = None) -> list[str]:
    """지금 팔리는 시즌 (3~8월 봄 · 여름, 그 외 가을 · 겨울 — 기획 · 사입 포함)"""
    m = (now or date.today()).month
    return ["C0071", "C0072", "C0075", "C0076", "C007A", "C007B"] if 3 <= m <= 8 else ["C0073", "C0074", "C0077", "C0078", "C007C", "C007D"]


def default_plan_years(now: date | None = None) -> list[str]:
    """지금 팔리는 기획년도 (1~2월은 지난해 가을 · 겨울이 함께 팔림)"""
    d = now or date.today()
    return [str(d.year - 1), str(d.year)] if d.month <= 2 else [str(d.year)]


def today() -> str:
    return datetime.now().strftime("%Y%m%d")


def ymd_label(v: str | None) -> str | None:
    return f"{v[:4]}-{v[4:6]}-{v[6:8]}" if v and len(v) >= 8 else None


# ---------------------------------------------------------------- 조건 확인
def brand_options(allowed: list[str] | None) -> list[dict]:
    """사용자가 볼 수 있는 브랜드 (브랜드 데이터 권한 = 브랜드명 목록, None 이면 전부)"""
    return [{"code": c, "name": n} for c, n in BRAND_CODES.items() if allowed is None or n in allowed]


def brand_code(brand: str | None, allowed: list[str] | None) -> str:
    opts = brand_options(allowed)
    if not opts:
        bad("볼 수 있는 브랜드가 없습니다. 관리자에게 브랜드 권한을 요청하세요.", 403)
    if not brand:
        return opts[0]["code"]
    b = brand.strip()
    code = b.upper() if b.upper() in BRAND_CODES else next((c for c, n in BRAND_CODES.items() if n == b), None)
    if not code:
        bad(f"브랜드는 {', '.join(o['name'] for o in opts)} 중 하나입니다.")
    if not any(o["code"] == code for o in opts):
        bad(f"'{BRAND_CODES[code]}' 브랜드 권한이 없습니다.", 403)
    return code


def _d(v: str) -> date:
    return datetime.strptime(v, "%Y%m%d").date()


def period(frm: str | None, to: str | None, default_days: int, *, end_yesterday: bool = False) -> tuple[str, str]:
    """판매 기간 (YYYYMMDD · YYYY-MM-DD). 최대 31일, 끝은 오늘 이후 불가. 비우면 최근 default_days 일"""
    def norm(v, key):
        if v is None or v == "":
            return None
        s = str(v).replace("-", "").strip()
        if not _D8.match(s):
            bad(f"{key} 는 YYYYMMDD 형식입니다.")
        try:
            _d(s)
        except ValueError:
            bad(f"{key} 날짜가 올바르지 않습니다.")
        return s
    t, f = norm(to, "기간 끝"), norm(frm, "기간 시작")
    now = date.today()
    end = _d(t) if t else (now - timedelta(days=1) if end_yesterday else now)
    if end > now:
        bad("기간 끝은 오늘 이후일 수 없습니다.")
    start = _d(f) if f else end - timedelta(days=default_days - 1)
    if start > end:
        bad("기간 시작이 끝보다 늦습니다.")
    if (end - start).days + 1 > MAX_DAYS:
        bad(f"기간은 최대 {MAX_DAYS}일(한 달)까지 조회합니다.")
    return start.strftime("%Y%m%d"), end.strftime("%Y%m%d")


def code_list(v, key: str, pattern: str = r"^[A-Za-z0-9_]{1,20}$", limit: int = 40) -> list[str]:
    if v is None or v == "":
        return []
    items = v.split(",") if isinstance(v, str) else v
    if not isinstance(items, list):
        bad(f"{key} 는 목록입니다.")
    out = [str(x).strip() for x in items if str(x).strip()]
    if len(out) > limit or any(not re.match(pattern, x) for x in out):
        bad(f"{key} 값이 올바르지 않습니다.")
    return sorted(set(out))


def prdt_prefix(v) -> str | None:
    if not v:
        return None
    s = str(v).strip().upper()
    if not re.match(r"^[A-Z0-9]{2,20}$", s):
        bad("품번은 영문 대문자 · 숫자 2~20자입니다 (앞부분만 넣어도 됩니다).")
    return s


def binds(values: list[str], prefix: str) -> tuple[str, dict]:
    b = {f"{prefix}{i}": v for i, v in enumerate(values)}
    return ", ".join(":" + k for k in b) or "NULL", b


def chunks(values: list, n: int = 500):
    for i in range(0, len(values), n):
        yield values[i:i + n]


# ---------------------------------------------------------------- 매장 기준 정보
# 행사 · 가상 매장 (재고 재배치 추천 · 통계에서 제외): ERP 매장 마스터의 가상매장구분(C708) · 유통형태(C002) · 매장MD유형(C387) · 매장관리구분(C377) · 실매장 여부
VIRTUAL_FORMS = {"C00206": "해외", "C00207": "특판", "C00208": "사내", "C00209": "기타 유통", "C00210": "사입/판매분자동생성"}
VIRTUAL_MD = {"C3876": "해외법인", "C3877": "홈쇼핑", "C3878": "이관", "C3879": "기타 MD", "C3880": "EBIZ", "C3881": "사내행사"}
VIRTUAL_MGN = {"C3775": "홈쇼핑", "C3776": "이관", "C3777": "행사", "C3778": "업체", "C3779": "기타 관리"}
VIRTUAL_SIMUL = {"C708010": "LOSS매장", "C708020": "재고조정매장", "C708030": "해외매장", "C708040": "사내매장", "C708050": "온라인매장",
                 "C708060": "클레임판매매장", "C708070": "오픈매장", "C708100": "직원구매매장", "C708110": "B품 매장"}


def virtual_reason(simul: str | None, form: str | None, md: str | None, mgn: str | None, real: str | None) -> str | None:
    """행사 · 가상 매장이면 이유(가장 구체적인 것 하나), 아니면 None"""
    if simul:
        return VIRTUAL_SIMUL.get(simul, "가상매장")
    if md in VIRTUAL_MD:
        return VIRTUAL_MD[md]
    if mgn in VIRTUAL_MGN:
        return VIRTUAL_MGN[mgn]
    if form in VIRTUAL_FORMS:
        return VIRTUAL_FORMS[form]
    if real == "N":
        return "실매장 아님"
    return None


def shops() -> dict[str, dict]:
    """매장 마스터 + 자동 RT 그룹 설정 (10분 캐시)"""
    def load():
        out: dict[str, dict] = {}
        for sid, nm, mo, dshop, team, typ, attr2, simul, form, md, mgn, real in db.query(
                """SELECT SHOP_ID, SHOP_NM, MO_BRD_CD, DSHOP_CLSBY, TEAM_CD, SHOP_TYPE, ATTR2, SIMUL_SHOP_CLSBY, SHOP_FORM, SHOP_MD_TYPE,
                          MGN_CLSBY, REAL_SHOP_YN FROM T_SHOP""")[1]:
            why = virtual_reason(simul, form, md, mgn, real)
            out[sid] = {"shopId": sid, "shopNm": nm, "moBrd": mo, "normal": dshop == "C0500", "team": team, "type": typ,
                        "attr2": attr2, "rt": None, "virtual": bool(why), "virtualWhy": why}
        for r in db.query("""SELECT SHOP_ID, RT_GRP_ID, REQ_ABLE_QTY, REQ_ABLE_YN, ASIGN_ABLE_QTY, MIN_RETAIN_STOCK_QTY,
                                    F_DELV_AF_DAYS, L_DELV_AF_DAYS
                               FROM T_SHOP_RT_GRP_DETL""")[1]:
            s = out.get(r[0])
            if s is None:
                continue
            s["rt"] = {"grp": r[1], "reqAble": None if r[2] is None else int(r[2]), "reqAbleYn": r[3], "asign": int(r[4] or 0), "minRetain": int(r[5] or 0),
                       "fDays": int(r[6] or 0), "lDays": int(r[7] or 0)}
        return out
    return cached(("shops",), CACHE_TTL, load)


def team_names() -> dict[str, str]:
    return cached(("teams",), CACHE_TTL, lambda: dict(db.query(
        "SELECT CD_KEY, CD_NM FROM T_COMN_CD WHERE PARENT_CD = 'C620' AND LANG_CLSBY = 'ko'")[1]))


def code_names(parent: str) -> dict[str, str]:
    return cached(("comn", parent), CACHE_TTL, lambda: dict(db.query(
        "SELECT CD_KEY, CD_NM FROM T_COMN_CD WHERE PARENT_CD = :p AND LANG_CLSBY = 'ko'", {"p": parent})[1]))


SHOP_TYPE_NM = {"C0041": "백화점", "C0042": "대리점", "C0043": "아울렛", "C0044": "직영점"}


def base_grade_group(brand: str) -> str | None:
    rows = db.query("""SELECT GRD_GRP_ID FROM T_SHOP_GRD_GRP_INFO
                        WHERE COMPY_CD = :c AND PARENT_BRD_CD = :b AND BASE_GRP_YN = 'Y'""", {"c": COMPY_CD, "b": brand})[1]
    return rows[0][0] if rows else None


def grade_shops(grp: str) -> dict[str, dict]:
    """등급 그룹의 매장 (매장등급이 있는 매장만) — 등급 우선순위는 공통코드 C094 PRTY_RANK"""
    def load():
        return {sid: {"grd": grd, "grdNm": nm, "prty": prty, "rank": rank} for sid, grd, nm, prty, rank in db.query(
            """SELECT D.SHOP_ID, D.SHOP_GRD, C.CD_NM, C.PRTY_RANK, D.SHOP_GRD_RANK
                 FROM T_SHOP_GRD_GRP_DETL D, T_COMN_CD C
                WHERE D.GRD_GRP_ID = :g AND D.SHOP_GRD IS NOT NULL
                  AND C.PARENT_CD(+) = 'C094' AND C.LANG_CLSBY(+) = 'ko' AND C.CD_KEY(+) = D.SHOP_GRD""", {"g": grp})[1]}
    return cached(("grade", grp), CACHE_TTL, load)


def style_info(prdts: list[str]) -> dict[str, dict]:
    """품번 → 스타일명 · 기획년도 · 시즌 · 아이템 · 상품구분"""
    out: dict[str, dict] = {}
    for part in chunks(sorted(set(prdts)), 900):
        ph, b = binds(part, "p")
        for p, nm, yy, sesn, item, goods, brd in db.query(
                f"""SELECT PRDT_CD, STYLE_NM, PLAN_YY, SESN_CD, ITEM_CD, GOODS_CLSBY, BRD_CD FROM T_STYLE_PLAN
                     WHERE COMPY_CD = '{COMPY_CD}' AND PRDT_CD IN ({ph})""", b)[1]:
            out[p] = {"styleNm": nm, "planYy": yy, "sesn": sesn, "item": item, "goods": goods, "brd": brd}
    return out


# ---------------------------------------------------------------- 수불제어
@dataclass
class Controls:
    """한 브랜드 · 한 날짜의 유효 수불제어 · 자동RT 제외 스타일"""
    flags: dict[str, dict] = field(default_factory=dict)                     # CNTR_ID → 플래그
    shop_cntrs: dict[str, set] = field(default_factory=dict)                 # 매장 → 매장 지정 제어 (삭제 안 된 매장 행)
    has_shop: set = field(default_factory=set)                               # 매장 행이 하나라도 있는 제어 (삭제 포함)
    has_prdt: set = field(default_factory=set)                               # 제품 행이 하나라도 있는 제어 (삭제 포함)
    prdt_all: dict[tuple, set] = field(default_factory=dict)                 # (제어, 품번) → 칼라 (삭제 포함)
    prdt_live: dict[str, list] = field(default_factory=dict)                 # 품번 → [(제어, 칼라)] (삭제 제외)
    xcld: list = field(default_factory=list)                                 # (팀, 년도, 시즌, 아이템, 품번)

    def _hit(self, cntrs, shop: str, prdt: str, color: str, f1, f2) -> bool:
        def color_ok(c):
            return c == "*" or c == color
        for cid in self.shop_cntrs.get(shop, ()):
            if cid in cntrs and f1(self.flags[cid]):
                if cid not in self.has_prdt or any(color_ok(c) for c in self.prdt_all.get((cid, prdt), ())):
                    return True
        for cid, c in self.prdt_live.get(prdt, ()):
            if cid in cntrs and color_ok(c) and cid not in self.has_shop and f2(self.flags[cid]):
                return True
        return False

    def controlled_id(self, shop: str, prdt: str, color: str, clsby: str) -> bool:
        """F_GET_RNDS_CNTR_ID(…, clsby) 가 제어 ID 를 돌려주면 True ('13' · '36')"""
        def f(cols):
            return lambda fl: "Y" in "".join(fl.get(c) or "" for c in cols)
        return self._hit(self.flags, shop, prdt, color, f(_FLAGS1[clsby]), f(_FLAGS2[clsby]))

    def controlled_rt_out(self, shop: str, prdt: str, color: str, shop_info: dict | None, style: dict | None) -> bool:
        """F_GET_RNDS_CNTR_AUTO_RT(…, '32') = 'Y' : 자동RT 제어 · 자동RT 제외 스타일(팀) · 매장 요청가능여부 N"""
        def rt(fl):
            return fl.get("AUTO_RT_CNTR_YN") == "Y"
        if self._hit(self.flags, shop, prdt, color, rt, rt):
            return True
        team = (shop_info or {}).get("team")
        if style and team:
            for t, yy, sesn, item, p in self.xcld:
                if t == team and yy == style.get("planYy") and sesn == style.get("sesn") \
                        and item in ("*", style.get("item")) and p in ("*", prdt):
                    return True
        return ((shop_info or {}).get("rt") or {}).get("reqAbleYn") != "Y"


def controls(brand: str, dt: str | None = None) -> Controls:
    dt = dt or today()

    def load() -> Controls:
        c = Controls()
        cols = ", ".join(_FLAG_COLS)
        for r in db.query(f"""SELECT CNTR_ID, {cols} FROM T_RNDS_CNTR
                               WHERE COMPY_CD = :c AND PARENT_BRD_CD = :b AND DEL_DAY IS NULL
                                 AND :dt BETWEEN APLY_DT AND NVL(CANCL_DT, END_DT)""", {"c": COMPY_CD, "b": brand, "dt": dt})[1]:
            c.flags[r[0]] = dict(zip(_FLAG_COLS, r[1:]))
        ids = sorted(c.flags)
        for part in chunks(ids, 500):
            ph, b = binds(part, "i")
            for cid, sid, deleted in db.query(f"SELECT CNTR_ID, SHOP_ID, DEL_DAY FROM T_RNDS_CNTR_SHOP WHERE CNTR_ID IN ({ph})", b)[1]:
                c.has_shop.add(cid)
                if deleted is None:
                    c.shop_cntrs.setdefault(sid, set()).add(cid)
            for cid, p, color, deleted in db.query(
                    f"SELECT CNTR_ID, PRDT_CD, COLOR_CD, DEL_DAY FROM T_RNDS_CNTR_PRDT WHERE CNTR_ID IN ({ph})", b)[1]:
                c.has_prdt.add(cid)
                c.prdt_all.setdefault((cid, p), set()).add(color)
                if deleted is None:
                    c.prdt_live.setdefault(p, []).append((cid, color))
        c.xcld = [tuple(r) for r in db.query(
            """SELECT A.TEAM_CD, B.YEAR_CD, B.SEASON_CD, B.ITEM_CD, B.PRDT_CD
                 FROM T_AUTO_RT_XCLD A, T_AUTO_RT_XCLD_PRDT B
                WHERE A.COMPY_CD = :c AND A.PARENT_BRD_CD = :b AND A.XCLD_ID = B.XCLD_ID AND A.PARENT_BRD_CD = B.PARENT_BRD_CD
                  AND A.COMPY_CD = B.COMPY_CD AND B.DEL_DAY IS NULL AND :dt BETWEEN A.APLY_START_DT AND A.APLY_END_DT""",
            {"c": COMPY_CD, "b": brand, "dt": dt})[1]]
        return c
    return cached(("ctl", brand, dt), CTL_TTL, load)
