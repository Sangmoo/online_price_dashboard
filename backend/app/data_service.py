"""온라인 가격 수집 데이터(T_SELECT_ONLINE_MNG_R) 조회/집계 서비스."""
from __future__ import annotations

import io
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from . import config, db

TABLE = "T_SELECT_ONLINE_MNG_R"

# 화면/엑셀 컬럼 정의 (키, 한글 라벨)
COLUMNS: list[tuple[str, str]] = [
    ("ONLINE_ID", "ONLINE_ID"),
    ("DT", "수집일"),
    ("PRDT_CD", "상품코드"),
    ("PRICE", "기준가"),
    ("DC_PRICE", "사이트_할인가"),
    ("DC_RATE", "할인율(%)"),
    ("MALL_NM", "사이트명"),
    ("TITLE", "TITLE"),
    ("RMK", "매장정보"),
    ("INS_DAY", "수집시간"),
    ("NAVER_PAY_SELL_NO", "판매자ID"),
    ("URL", "URL"),
]
COLUMN_KEYS = [c[0] for c in COLUMNS]
NUMERIC_COLS = {"ONLINE_ID", "PRICE", "DC_PRICE", "DC_RATE"}

# 할인율(%) 계산식 — 기준가 대비 사이트 할인가
DC_RATE_SQL = "ROUND((PRICE - DC_PRICE) / NULLIF(PRICE, 0) * 100, 2)"


# ----------------------------------------------------------------------------
# 공통 유틸
# ----------------------------------------------------------------------------
def parse_dt(value: str) -> date:
    try:
        return datetime.strptime(value.replace("-", ""), "%Y%m%d").date()
    except (ValueError, AttributeError):
        raise ValueError(f"날짜 형식이 올바르지 않습니다: {value!r} (YYYYMMDD)")


def fmt_dt(d: date) -> str:
    return d.strftime("%Y%m%d")


def validate_range(start: str, end: str) -> tuple[str, str]:
    s, e = parse_dt(start), parse_dt(end)
    if s > e:
        s, e = e, s
    if (e - s).days + 1 > config.MAX_RANGE_DAYS:
        raise ValueError(f"조회 기간은 최대 {config.MAX_RANGE_DAYS}일까지 가능합니다.")
    return fmt_dt(s), fmt_dt(e)


def _num(v):
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return v
    return int(f) if f.is_integer() else round(f, 2)


class _TTLCache:
    def __init__(self, max_items: int):
        self.max_items = max_items
        self._data: dict = {}
        self._lock = threading.Lock()

    def get(self, key):
        with self._lock:
            item = self._data.get(key)
            if item and item[0] > time.time():
                return item[1]
            self._data.pop(key, None)
            return None

    def set(self, key, value, ttl: float):
        with self._lock:
            if len(self._data) >= self.max_items:
                oldest = min(self._data, key=lambda k: self._data[k][0])
                self._data.pop(oldest, None)
            self._data[key] = (time.time() + ttl, value)


def _ttl_for(dt_to: str) -> int:
    # 오늘 데이터는 수집 중이므로 짧게 캐시
    return 300 if dt_to >= fmt_dt(date.today()) else 3600


_day_cache = _TTLCache(max_items=6)
_dash_cache = _TTLCache(max_items=20)


# ----------------------------------------------------------------------------
# 수집일 목록
# ----------------------------------------------------------------------------
def available_dates(days_back: int = 120) -> list[dict]:
    since = fmt_dt(date.today() - timedelta(days=days_back))
    rows = db.query_dicts(
        f"SELECT DT, COUNT(*) AS CNT FROM {TABLE} WHERE DT >= :since GROUP BY DT ORDER BY DT DESC",
        {"since": since},
    )
    return [{"dt": r["DT"], "count": int(r["CNT"])} for r in rows]


# ----------------------------------------------------------------------------
# 일자별 상세 (그리드 / 엑셀)
# ----------------------------------------------------------------------------
def load_day(dt: str) -> pd.DataFrame:
    dt = fmt_dt(parse_dt(dt))
    cached = _day_cache.get(dt)
    if cached is not None:
        return cached
    cols, rows = db.query(
        f"""
        SELECT ONLINE_ID, DT, PRDT_CD, PRICE, DC_PRICE, {DC_RATE_SQL} AS DC_RATE,
               MALL_NM, TITLE, RMK, INS_DAY, NAVER_PAY_SELL_NO, URL
          FROM {TABLE}
         WHERE DT = :dt
        """,
        {"dt": dt},
    )
    df = pd.DataFrame(rows, columns=cols)
    for c in ("PRICE", "DC_PRICE", "DC_RATE", "ONLINE_ID"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    _day_cache.set(dt, df, _ttl_for(dt))
    return df


def _filter_sort(df: pd.DataFrame, q: str | None, mall: str | None, sort: str | None, order: str,
                 min_rate: float | None = None, max_rate: float | None = None) -> pd.DataFrame:
    if min_rate is not None:
        df = df[df["DC_RATE"] >= min_rate]
    if max_rate is not None:
        df = df[df["DC_RATE"] <= max_rate]
    if mall:
        df = df[df["MALL_NM"] == mall]
    if q:
        q = q.strip().lower()
        mask = pd.Series(False, index=df.index)
        for c in ("PRDT_CD", "TITLE", "MALL_NM", "RMK", "NAVER_PAY_SELL_NO"):
            mask |= df[c].fillna("").astype(str).str.lower().str.contains(q, regex=False)
        df = df[mask]
    if sort in COLUMN_KEYS:
        df = df.sort_values(sort, ascending=(order != "desc"), na_position="last", kind="stable")
    return df


def day_rows(dt: str, page: int, size: int, sort: str | None, order: str, q: str | None, mall: str | None,
             min_rate: float | None = None, max_rate: float | None = None) -> dict:
    base = load_day(dt)
    df = _filter_sort(base, q, mall, sort, order, min_rate, max_rate)
    size = max(10, min(size, 500))
    total = len(df)
    pages = max(1, -(-total // size))
    page = max(1, min(page, pages))
    chunk = df.iloc[(page - 1) * size : page * size]
    records = [
        {k: (_num(v) if k in NUMERIC_COLS else v) for k, v in rec.items()}
        for rec in chunk.astype(object).where(chunk.notna(), None).to_dict("records")
    ]
    return {
        "dt": fmt_dt(parse_dt(dt)),
        "columns": [{"key": k, "label": l} for k, l in COLUMNS],
        "rows": records,
        "total": total,
        "totalAll": len(base),
        "page": page,
        "pages": pages,
        "size": size,
        "malls": sorted(base["MALL_NM"].dropna().unique().tolist()),
        "summary": {
            "products": int(df["PRDT_CD"].nunique()),
            "malls": int(df["MALL_NM"].nunique()),
            "avgDcRate": _num(df["DC_RATE"].mean()) if total else None,
        },
    }


HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(bold=True, color="FFFFFF")
COL_WIDTHS = {"ONLINE_ID": 11, "DT": 11, "PRDT_CD": 16, "PRICE": 11, "DC_PRICE": 13, "DC_RATE": 10,
              "MALL_NM": 16, "TITLE": 45, "RMK": 50, "INS_DAY": 15, "NAVER_PAY_SELL_NO": 16, "URL": 60}


def write_xlsx(sheet_title: str, columns: list[tuple[str, str]], rows) -> bytes:
    """(키, 라벨) 컬럼 정의와 dict 로우 이터러블로 xlsx 생성."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_title[:31])
    ws.freeze_panes = "A2"
    for i, (k, _) in enumerate(columns):
        ws.column_dimensions[_col_letter(i)].width = COL_WIDTHS.get(k, 14)
    header = []
    for _, label in columns:
        cell = WriteOnlyCell(ws, value=label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        header.append(cell)
    ws.append(header)
    for r in rows:
        ws.append([_xl_value(r.get(k)) for k, _ in columns])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xl_value(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    return v


def _col_letter(idx: int) -> str:
    s, idx = "", idx + 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


def export_day(dt: str, sort: str | None, order: str, q: str | None, mall: str | None,
               min_rate: float | None = None, max_rate: float | None = None, cols: list[str] | None = None) -> bytes:
    df = _filter_sort(load_day(dt), q, mall, sort, order, min_rate, max_rate)
    df = df.astype(object).where(df.notna(), None)
    columns = [c for c in COLUMNS if not cols or c[0] in cols] or COLUMNS
    return write_xlsx(f"온라인가격_{dt}", columns, df.to_dict("records"))


# ----------------------------------------------------------------------------
# 대시보드 (기간 집계, 최대 31일)
# ----------------------------------------------------------------------------
def dashboard(start: str, end: str) -> dict:
    start, end = validate_range(start, end)
    key = (start, end)
    cached = _dash_cache.get(key)
    if cached is not None:
        return cached

    p = {"s": start, "e": end}
    where = f"FROM {TABLE} WHERE DT BETWEEN :s AND :e"

    jobs = {}
    with ThreadPoolExecutor(max_workers=6) as pool:
        jobs["kpi"] = pool.submit(
            db.query_dicts,
            f"""
            SELECT COUNT(*) AS ROW_CNT, COUNT(DISTINCT PRDT_CD) AS PRDT_CNT,
                   COUNT(DISTINCT MALL_NM) AS MALL_CNT, COUNT(DISTINCT NAVER_PAY_SELL_NO) AS SELLER_CNT,
                   COUNT(DISTINCT DT) AS DAY_CNT,
                   ROUND(AVG({DC_RATE_SQL}), 2) AS AVG_DC_RATE,
                   MAX({DC_RATE_SQL}) AS MAX_DC_RATE,
                   SUM(CASE WHEN {DC_RATE_SQL} >= 30 THEN 1 ELSE 0 END) AS DEEP_DC_CNT
            {where}
            """,
            p,
        )

        jobs["daily"] = pool.submit(
            db.query_dicts,
            f"""
            SELECT DT, COUNT(*) AS ROW_CNT, COUNT(DISTINCT PRDT_CD) AS PRDT_CNT,
                   COUNT(DISTINCT MALL_NM) AS MALL_CNT, ROUND(AVG({DC_RATE_SQL}), 2) AS AVG_DC_RATE
            {where}
            GROUP BY DT ORDER BY DT
            """,
            p,
        )

        jobs["malls"] = pool.submit(
            db.query_dicts,
            f"""
            SELECT * FROM (
                SELECT MALL_NM, COUNT(*) AS ROW_CNT, COUNT(DISTINCT PRDT_CD) AS PRDT_CNT,
                       ROUND(AVG({DC_RATE_SQL}), 2) AS AVG_DC_RATE
                {where}
                GROUP BY MALL_NM ORDER BY COUNT(*) DESC
            ) WHERE ROWNUM <= 15
            """,
            p,
        )

        jobs["mall_dc"] = pool.submit(
            db.query_dicts,
            f"""
            SELECT * FROM (
                SELECT MALL_NM, COUNT(*) AS ROW_CNT, ROUND(AVG({DC_RATE_SQL}), 2) AS AVG_DC_RATE
                {where}
                GROUP BY MALL_NM HAVING COUNT(*) >= 50
                ORDER BY AVG({DC_RATE_SQL}) DESC
            ) WHERE ROWNUM <= 10
            """,
            p,
        )

        jobs["hist"] = pool.submit(
            db.query_dicts,
            f"""
            SELECT LEAST(FLOOR({DC_RATE_SQL} / 5) * 5, 70) AS BUCKET, COUNT(*) AS ROW_CNT
            {where} AND PRICE > 0
            GROUP BY LEAST(FLOOR({DC_RATE_SQL} / 5) * 5, 70) ORDER BY 1
            """,
            p,
        )

        jobs["products"] = pool.submit(
            db.query_dicts,
            f"""
            SELECT * FROM (
                SELECT * FROM (
                    SELECT PRDT_CD,
                           MAX(TITLE) KEEP (DENSE_RANK FIRST ORDER BY DC_PRICE) AS TITLE,
                           MAX(PRICE) AS PRICE,
                           MIN(DC_PRICE) AS MIN_DC_PRICE,
                           MAX(DC_RATE) AS MAX_DC_RATE,
                           MAX(MALL_NM) KEEP (DENSE_RANK FIRST ORDER BY DC_PRICE) AS MIN_MALL,
                           MAX(DT) KEEP (DENSE_RANK FIRST ORDER BY DC_PRICE) AS MIN_DT,
                           COUNT(DISTINCT MALL_NM) AS MALL_CNT,
                           COUNT(*) AS ROW_CNT
                      FROM (SELECT PRDT_CD, TITLE, PRICE, DC_PRICE, MALL_NM, DT, {DC_RATE_SQL} AS DC_RATE
                            {where} AND PRICE > 0)
                     GROUP BY PRDT_CD
                ) ORDER BY MAX_DC_RATE DESC NULLS LAST, ROW_CNT DESC
            ) WHERE ROWNUM <= 20
            """,
            p,
        )
    kpi = jobs["kpi"].result()[0]
    daily, malls, mall_dc, hist, products = (
        jobs[k].result() for k in ("daily", "malls", "mall_dc", "hist", "products")
    )

    def clean(rows):
        return [{k: _num(v) if not isinstance(v, str) else v for k, v in r.items()} for r in rows]

    result = {
        "start": start,
        "end": end,
        "kpi": clean([kpi])[0],
        "daily": clean(daily),
        "malls": clean(malls),
        "mallDiscount": clean(mall_dc),
        "histogram": clean(hist),
        "topProducts": clean(products),
        "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    _dash_cache.set(key, result, _ttl_for(end))
    return result


def product_online(prdt_cd: str, days: int = 31) -> dict:
    """상품 팝업의 온라인 가격: 최근 수집일별 평균 할인율 · 최저 할인가 · 정상가 · 사이트 수, 마지막 수집일의 사이트별 최저가.
    인덱스 (PRDT_CD, DT) 로 바로 찾는다."""
    since = (datetime.now() - timedelta(days=days)).strftime("%Y%m%d")
    daily = db.query_dicts(
        f"""SELECT DT, ROUND(AVG({DC_RATE_SQL}), 1) AS AVG_DC_RATE, MAX({DC_RATE_SQL}) AS MAX_DC_RATE, MIN(DC_PRICE) AS MIN_DC_PRICE,
                   MAX(PRICE) AS LIST_PRICE, COUNT(DISTINCT MALL_NM) AS MALL_CNT, COUNT(*) AS ROW_CNT
              FROM {TABLE} WHERE PRDT_CD = :p AND DT >= :s AND PRICE > 0 GROUP BY DT ORDER BY DT""",
        {"p": prdt_cd, "s": since})
    last = daily[-1]["DT"] if daily else None
    malls, title = [], None
    if last:
        rows = db.query_dicts(
            f"""SELECT MALL_NM, MIN(DC_PRICE) AS LOW_PRICE, MAX(PRICE) AS LIST_PRICE, MAX({DC_RATE_SQL}) AS TOP_RATE, MAX(TITLE) AS ANY_TITLE
                  FROM {TABLE} WHERE PRDT_CD = :p AND DT = :d AND PRICE > 0 GROUP BY MALL_NM ORDER BY 2""",
            {"p": prdt_cd, "d": last})
        malls = [{"mallNm": r["MALL_NM"], "dcPrice": _num(r["LOW_PRICE"]), "price": _num(r["LIST_PRICE"]), "dcRate": _num(r["TOP_RATE"])}
                 for r in rows[:15]]
        title = next((r["ANY_TITLE"] for r in rows if r["ANY_TITLE"]), None)
    return {"title": title, "lastDt": last,
            "daily": [{"dt": r["DT"], "avgDcRate": _num(r["AVG_DC_RATE"]), "maxDcRate": _num(r["MAX_DC_RATE"]),
                       "minDcPrice": _num(r["MIN_DC_PRICE"]), "price": _num(r["LIST_PRICE"]), "malls": int(r["MALL_CNT"]),
                       "rows": int(r["ROW_CNT"])} for r in daily],
            "malls": malls}
